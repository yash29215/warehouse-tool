
import customtkinter as ctk
from tkinter import filedialog, messagebox
import tkinter as tk
from tkinter import ttk
import json
import pandas as pd
import os
import math
import numpy as np
from collections import defaultdict, deque
import heapq
from openpyxl import Workbook, load_workbook

# ─────────────────────────────────────────────
# Theme configuration
# ─────────────────────────────────────────────
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

ACCENT     = "#2563EB"       # vivid blue
ACCENT_ALT = "#3B82F6"       # lighter blue for hover
SUCCESS    = "#10B981"
WARNING    = "#F59E0B"
ERROR      = "#EF4444"

BG_DARK    = "#0F172A"       # slate-900
BG_PANEL   = "#1E293B"       # slate-800
BG_CARD    = "#1E293B"
BG_INPUT   = "#0F172A"
BORDER     = "#334155"       # slate-700
TEXT_HI    = "#F1F5F9"       # slate-100
TEXT_LO    = "#94A3B8"       # slate-400
TEXT_MID   = "#CBD5E1"       # slate-300

FONT_TITLE = ("Segoe UI", 22, "bold")
FONT_HEAD  = ("Segoe UI", 12, "bold")
FONT_LABEL = ("Segoe UI", 11)
FONT_SMALL = ("Segoe UI", 9)
FONT_MONO  = ("Consolas", 10)


# ─────────────────────────────────────────────
# Small helper widgets
# ─────────────────────────────────────────────

def make_section_label(parent, text):
    """Styled section heading."""
    lbl = ctk.CTkLabel(parent, text=text, font=FONT_HEAD,
                        text_color=TEXT_LO, anchor="w")
    return lbl


def make_separator(parent):
    frm = ctk.CTkFrame(parent, fg_color=BORDER, height=1, corner_radius=0)
    return frm


class FilePickerRow(ctk.CTkFrame):
    """A labelled file-picker row (label + entry + button)."""
    def __init__(self, parent, label: str, filetypes, on_browse=None,
                 default_text="", **kwargs):
        super().__init__(parent, fg_color="transparent", **kwargs)
        self._on_browse = on_browse
        self._filetypes = filetypes

        ctk.CTkLabel(self, text=label, font=FONT_LABEL,
                     text_color=TEXT_MID, width=130, anchor="w").pack(side="left")
        self.entry = ctk.CTkEntry(self, placeholder_text=default_text,
                                  fg_color=BG_INPUT, border_color=BORDER,
                                  text_color=TEXT_HI, height=36,
                                  corner_radius=8, font=FONT_LABEL)
        self.entry.pack(side="left", fill="x", expand=True, padx=(8, 8))
        self.btn = ctk.CTkButton(self, text="Browse", width=90,
                                 fg_color=BG_PANEL, hover_color=BORDER,
                                 border_color=BORDER, border_width=1,
                                 text_color=TEXT_MID, corner_radius=8,
                                 height=36, font=FONT_LABEL,
                                 command=self._browse)
        self.btn.pack(side="left")

    def _browse(self):
        path = filedialog.askopenfilename(filetypes=self._filetypes)
        if path:
            self.entry.delete(0, "end")
            self.entry.insert(0, path)
            if self._on_browse:
                self._on_browse(path)

    def get(self):
        return self.entry.get()

    def set(self, val):
        self.entry.delete(0, "end")
        self.entry.insert(0, val)


class LogBox(ctk.CTkTextbox):
    """Scrollable mono log output."""
    def __init__(self, parent, **kwargs):
        super().__init__(parent, font=FONT_MONO,
                         fg_color=BG_INPUT, text_color=TEXT_MID,
                         border_color=BORDER, border_width=1,
                         corner_radius=8, wrap="word", **kwargs)
        self.configure(state="normal")

    def log(self, msg: str, tag: str = ""):
        self.configure(state="normal")
        prefix = {"ok": "✅ ", "err": "❌ ", "warn": "⚠️  ", "info": "ℹ️  "}.get(tag, "")
        self.insert("end", prefix + msg + "\n")
        self.see("end")
        self.configure(state="normal")
        self.update()

    def clear(self):
        self.configure(state="normal")
        self.delete("1.0", "end")


class AccentButton(ctk.CTkButton):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, fg_color=ACCENT, hover_color=ACCENT_ALT,
                         text_color="#FFFFFF", corner_radius=8,
                         height=38, font=("Segoe UI", 11, "bold"), **kwargs)


class GhostButton(ctk.CTkButton):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, fg_color=BG_PANEL, hover_color=BORDER,
                         border_color=BORDER, border_width=1,
                         text_color=TEXT_MID, corner_radius=8,
                         height=34, font=FONT_LABEL, **kwargs)


class Card(ctk.CTkFrame):
    def __init__(self, parent, title="", **kwargs):
        super().__init__(parent, fg_color=BG_PANEL,
                         corner_radius=12, border_color=BORDER,
                         border_width=1, **kwargs)
        if title:
            ctk.CTkLabel(self, text=title, font=FONT_HEAD,
                         text_color=TEXT_HI, anchor="w").pack(
                anchor="w", padx=16, pady=(14, 0))
            make_separator(self).pack(fill="x", padx=16, pady=(10, 0))


# ─────────────────────────────────────────────
# Styled Treeview (ttk — customtkinter has no native treeview)
# ─────────────────────────────────────────────

def make_styled_tree(parent, columns, headings, col_widths):
    style = ttk.Style()
    style.theme_use("clam")
    style.configure("Dark.Treeview",
                     background=BG_INPUT,
                     foreground=TEXT_MID,
                     fieldbackground=BG_INPUT,
                     rowheight=28,
                     borderwidth=0,
                     font=("Segoe UI", 10))
    style.configure("Dark.Treeview.Heading",
                     background=BG_PANEL,
                     foreground=TEXT_LO,
                     font=("Segoe UI", 10, "bold"),
                     relief="flat")
    style.map("Dark.Treeview",
              background=[("selected", ACCENT)],
              foreground=[("selected", "#FFFFFF")])

    frame = tk.Frame(parent, bg=BG_INPUT, bd=1, relief="solid",
                     highlightbackground=BORDER, highlightthickness=1)

    tree = ttk.Treeview(frame, columns=columns, show="headings",
                         style="Dark.Treeview")
    vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=vsb.set)

    for col, heading, width in zip(columns, headings, col_widths):
        tree.heading(col, text=heading)
        tree.column(col, width=width, anchor="center")

    tree.grid(row=0, column=0, sticky="nsew")
    vsb.grid(row=0, column=1, sticky="ns")
    frame.grid_rowconfigure(0, weight=1)
    frame.grid_columnconfigure(0, weight=1)
    return frame, tree


# ─────────────────────────────────────────────
# MAIN APPLICATION
# ─────────────────────────────────────────────

class WarehouseToolsApp:
    def __init__(self, root: ctk.CTk):
        self.root = root
        self.root.title("FILES AUTOMATION TOOL")
        self.root.geometry("1440x900")
        self.root.configure(fg_color=BG_DARK)
        self.root.minsize(1100, 720)

        # ── shared state ──
        self.json_data      = None
        self.aisles         = []
        self.loading_point  = None
        self.graph          = {}
        self.node_positions = {}
        self.excluded_aps   = []

        # N-deep feature state (BooleanVar/StringVar must be created
        # after the Tk root exists; they are initialised in setup_tab3)
        self.ndeep_enabled_var      = None  # ctk.BooleanVar
        self.ndeep_drop_seq_var     = None  # tk.StringVar  "f2l" | "l2f"
        self.ndeep_zone_scope_var   = None  # tk.StringVar  "BOTH"|"ENTRY"|"EXIT"
        self.ndeep_sub_scope_var    = None  # tk.StringVar  "BOTH"|"ENTRY"|"EXIT"
        self.ndeep_loc_scope_var    = None  # tk.StringVar  "BOTH"|"ENTRY"|"EXIT"
        self._ndeep_options_frame   = None  # ctk.CTkFrame revealed on toggle

        self._build_layout()

    # ─────────────────── Layout skeleton ───────────────────

    def _build_layout(self):
        # Top header
        self._build_header()

        # Left sidebar navigation
        self._sidebar = ctk.CTkFrame(self.root, fg_color=BG_PANEL,
                               corner_radius=0, width=210,
                               border_color=BORDER, border_width=0)
        self._sidebar.pack(side="left", fill="y", padx=0, pady=0)
        self._sidebar.pack_propagate(False)
        self._build_sidebar(self._sidebar)

        # Draggable resize handle between sidebar and content
        self._sash = tk.Frame(self.root, bg=BORDER, width=5, cursor="sb_h_double_arrow")
        self._sash.pack(side="left", fill="y")
        self._sash.bind("<Enter>",        lambda _: self._sash.configure(bg=ACCENT))
        self._sash.bind("<Leave>",        lambda _: self._sash.configure(bg=BORDER))
        self._sash.bind("<ButtonPress-1>",   self._sash_start)
        self._sash.bind("<B1-Motion>",       self._sash_drag)

        # Right content area
        self.content = ctk.CTkFrame(self.root, fg_color=BG_DARK,
                                    corner_radius=0)
        self.content.pack(side="left", fill="both", expand=True)

        # Pages (one frame each, stacked)
        self.pages = {}
        for name in ("tab1", "tab2", "tab3", "tab4"):
            frm = ctk.CTkFrame(self.content, fg_color=BG_DARK,
                                corner_radius=0)
            frm.place(relx=0, rely=0, relwidth=1, relheight=1)
            self.pages[name] = frm

        self.setup_tab1()
        self.setup_tab2()
        self.setup_tab3()
        self.setup_tab4()
        self._show_page("tab1")

        # Status bar
        self._build_statusbar()

    def _build_header(self):
        hdr = ctk.CTkFrame(self.root, fg_color=BG_PANEL,
                            corner_radius=0, height=58,
                            border_color=BORDER, border_width=0)
        hdr.pack(fill="x", side="top")
        hdr.pack_propagate(False)

        # left: logo text
        ctk.CTkLabel(hdr, text="⬡  FILES AUTOMATION TOOL",
                     font=("Segoe UI", 16, "bold"),
                     text_color=TEXT_HI).pack(side="left", padx=24, pady=0)

        # right: brand
        ctk.CTkLabel(hdr, text="GreyOrange",
                     font=("Segoe UI", 11),
                     text_color=TEXT_LO).pack(side="right", padx=24)

        # thin accent line at bottom of header
        line = ctk.CTkFrame(self.root, fg_color=ACCENT,
                             height=2, corner_radius=0)
        line.pack(fill="x", side="top")

    def _build_sidebar(self, sidebar):
        ctk.CTkLabel(sidebar, text="TOOLS", font=("Segoe UI", 9, "bold"),
                     text_color=TEXT_LO).pack(anchor="w", padx=20, pady=(24, 8))

        self._nav_buttons = {}
        nav_items = [
            ("tab1", "📊", "Each Pick"),
            ("tab2", "🔄", "Case Pick"),
            ("tab3", "🗺️", "Point to Point"),
            ("tab4", "✅", "Map Validator"),
        ]
        for key, icon, label in nav_items:
            btn = ctk.CTkButton(
                sidebar,
                text=f"  {icon}  {label}",
                anchor="w",
                fg_color="transparent",
                hover_color=BORDER,
                text_color=TEXT_MID,
                font=("Segoe UI", 11),
                height=40,
                corner_radius=8,
                command=lambda k=key: self._show_page(k)
            )
            btn.pack(fill="x", padx=10, pady=2)
            self._nav_buttons[key] = btn

    def _show_page(self, key: str):
        for k, frm in self.pages.items():
            frm.lower()
        self.pages[key].lift()

        # Highlight active nav button
        for k, btn in self._nav_buttons.items():
            if k == key:
                btn.configure(fg_color=ACCENT, text_color="#FFFFFF")
            else:
                btn.configure(fg_color="transparent", text_color=TEXT_MID)

    def _build_statusbar(self):
        bar = ctk.CTkFrame(self.root, fg_color=BG_PANEL,
                            corner_radius=0, height=30)
        bar.pack(fill="x", side="bottom")
        bar.pack_propagate(False)
        self.status_var = tk.StringVar(value="Ready")
        ctk.CTkLabel(bar, textvariable=self.status_var,
                     font=FONT_SMALL, text_color=TEXT_LO).pack(
            side="left", padx=16, pady=0)

    def update_status(self, msg: str):
        self.status_var.set(msg)
        self.root.update()

    # ─────────────────── Sidebar resize ───────────────────

    def _sash_start(self, event):
        self._sash_x_start = event.x_root
        self._sash_width_start = self._sidebar.winfo_width()

    def _sash_drag(self, event):
        delta = event.x_root - self._sash_x_start
        new_width = max(160, min(400, self._sash_width_start + delta))
        self._sidebar.configure(width=new_width)

    def _tab2_browse_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("Map files", "*.json *.smap"),
                       ("JSON", "*.json"), ("SMAP", "*.smap")])
        if path:
            self._tab2_file_entry.delete(0, "end")
            self._tab2_file_entry.insert(0, path)
            self._map_load_from_path(path)
            self.show_json_format_info()

    def _tab2_browse_template(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            self._tab2_template_entry.delete(0, "end")
            self._tab2_template_entry.insert(0, path)

    def _map_load_from_path(self, path):
        if not path or not os.path.exists(path):
            return
        try:
            with open(path, "r", encoding="utf-8") as fh:
                data = json.load(fh)
            self._map_points = []
            for pt in data.get("advancedPointList", []):
                cls = pt.get("className", "")
                pos = pt.get("pos", {})
                if "x" not in pos or "y" not in pos:
                    continue
                if cls in ("ActionPoint", "LocationMark"):
                    self._map_points.append({
                        "name": pt.get("instanceName", ""),
                        "x":    float(pos["x"]),
                        "y":    float(pos["y"]),
                        "type": "AP" if cls == "ActionPoint" else "LM",
                    })
            if not self._map_points:
                return
            self._map_excluded.clear()
            self._map_loading_aps.clear()
            self._map_unloading_aps.clear()
            self._map_cross_aisle_lms.clear()
            self._map_update_badges()
            self._map_fitted = False
            self._map_fit_all()
            self._map_redraw()
        except Exception as e:
            self.log_tab2(f"⚠️ Map load error: {e}")

    def _toggle_dup_options(self):
        if self.dup_enabled_var.get():
            self._dup_options_frame.pack(fill="x")
        else:
            self._dup_options_frame.pack_forget()

    # ─────────────────── Tab 1: Marker Extraction ───────────────────

    def setup_tab1(self):
        p = self.pages["tab1"]
        outer = ctk.CTkScrollableFrame(p, fg_color="transparent",
                                        scrollbar_button_color=BORDER,
                                        scrollbar_button_hover_color=ACCENT)
        outer.pack(fill="both", expand=True, padx=24, pady=20)

        # ── Page title ──
        ctk.CTkLabel(outer, text="Each Pick",
                     font=FONT_TITLE, text_color=TEXT_HI,
                     anchor="w").pack(anchor="w", pady=(0, 18))

        # ── Input file card ──
        card1 = Card(outer, title="Input File")
        card1.pack(fill="x", pady=(0, 14))
        inner1 = ctk.CTkFrame(card1, fg_color="transparent")
        inner1.pack(fill="x", padx=16, pady=(10, 16))

        self._tab1_file_picker = FilePickerRow(
            inner1, "JSON File", [("JSON", "*.json")],
            on_browse=self._on_tab1_browse)
        self._tab1_file_picker.pack(fill="x")

        # ── Marker types card ──
        card2 = Card(outer, title="Marker Types")
        card2.pack(fill="x", pady=(0, 14))
        inner2 = ctk.CTkFrame(card2, fg_color="transparent")
        inner2.pack(fill="x", padx=16, pady=(10, 16))

        chk_row = ctk.CTkFrame(inner2, fg_color="transparent")
        chk_row.pack(anchor="w")

        self.work_marker_var     = ctk.BooleanVar(value=True)
        self.resource_marker_var = ctk.BooleanVar(value=True)

        ctk.CTkCheckBox(chk_row, text="WorkMarker",
                        variable=self.work_marker_var,
                        font=FONT_LABEL, text_color=TEXT_MID,
                        fg_color=ACCENT, hover_color=ACCENT_ALT,
                        checkmark_color="#fff",
                        border_color=BORDER).pack(side="left", padx=(0, 24))
        ctk.CTkCheckBox(chk_row, text="ResourceMarker",
                        variable=self.resource_marker_var,
                        font=FONT_LABEL, text_color=TEXT_MID,
                        fg_color=ACCENT, hover_color=ACCENT_ALT,
                        checkmark_color="#fff",
                        border_color=BORDER).pack(side="left")

        # ── Output card ──
        card3 = Card(outer, title="Output")
        card3.pack(fill="x", pady=(0, 14))
        inner3 = ctk.CTkFrame(card3, fg_color="transparent")
        inner3.pack(fill="x", padx=16, pady=(10, 16))

        out_row = ctk.CTkFrame(inner3, fg_color="transparent")
        out_row.pack(fill="x")
        ctk.CTkLabel(out_row, text="Output File", font=FONT_LABEL,
                     text_color=TEXT_MID, width=130, anchor="w").pack(side="left")
        self.tab1_output_entry = ctk.CTkEntry(
            out_row, placeholder_text="workmarkers_output.xlsx",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            height=36, corner_radius=8, font=FONT_LABEL)
        self.tab1_output_entry.insert(0, "workmarkers_output.xlsx")
        self.tab1_output_entry.pack(side="left", fill="x", expand=True)

        # ── Run button ──
        AccentButton(outer, text="🚀  Extract Markers",
                     command=self.run_marker_extraction,
                     width=200).pack(anchor="w", pady=(6, 18))

        # ── Log card ──
        log_card = Card(outer, title="Log")
        log_card.pack(fill="both", expand=True)
        self.tab1_log = LogBox(log_card, height=220)
        self.tab1_log.pack(fill="both", expand=True, padx=16, pady=(10, 16))

    def _on_tab1_browse(self, path):
        self.show_marker_json_info()

    def browse_tab1_file(self):
        path = filedialog.askopenfilename(filetypes=[("JSON", "*.json")])
        if path:
            self._tab1_file_picker.set(path)
            self.show_marker_json_info()

    def log_tab1(self, msg: str):
        tag = "ok" if "✅" in msg else "err" if "❌" in msg or "💥" in msg else \
              "warn" if "⚠️" in msg else "info"
        clean = msg.lstrip("✅❌⚠️💥🔄🔍").strip()
        self.tab1_log.log(clean, tag)

    # ─────────────────── Tab 2: Pick Sequence ───────────────────

    def setup_tab2(self):
        p = self.pages["tab2"]

        # Single outer scroll — both sections scroll together
        outer = ctk.CTkScrollableFrame(
            p, fg_color="transparent",
            scrollbar_button_color=BORDER,
            scrollbar_button_hover_color=ACCENT)
        outer.pack(fill="both", expand=True)

        # ══ SECTION 1 — Input & Configuration ══
        sec1 = ctk.CTkFrame(outer, fg_color=BG_PANEL, corner_radius=12,
                             border_color=BORDER, border_width=1)
        sec1.pack(fill="x", padx=20, pady=(16, 0))

        # Section 1 header
        s1_hdr = ctk.CTkFrame(sec1, fg_color="transparent")
        s1_hdr.pack(fill="x", padx=20, pady=(14, 0))
        ctk.CTkLabel(s1_hdr, text="1", width=22, height=22,
                     fg_color=ACCENT, text_color="#fff", corner_radius=11,
                     font=("Segoe UI", 10, "bold")).pack(side="left", padx=(0, 10))
        _t = ctk.CTkFrame(s1_hdr, fg_color="transparent")
        _t.pack(side="left")
        ctk.CTkLabel(_t, text="Input & Configuration",
                     font=("Segoe UI", 14, "bold"), text_color=TEXT_HI,
                     anchor="w").pack(anchor="w")
        ctk.CTkLabel(_t, text="Upload map file and configure pick parameters",
                     font=FONT_SMALL, text_color=TEXT_LO, anchor="w").pack(anchor="w")

        # 4-card row
        crd_row = ctk.CTkFrame(sec1, fg_color="transparent")
        crd_row.pack(fill="x", padx=16, pady=(12, 0))
        c1 = Card(crd_row)
        c2 = Card(crd_row)
        c4 = Card(crd_row)
        for i, c in enumerate([c1, c2, c4]):
            c.pack(side="left", expand=True, fill="both",
                   padx=(0, 8) if i < 2 else 0)

        # ── Card 1: Map File ──
        ctk.CTkLabel(c1, text="MAP FILE", font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(
            anchor="w", padx=14, pady=(12, 6))
        dz = ctk.CTkFrame(c1, fg_color="#111827", corner_radius=8,
                          border_width=1, border_color="#203060")
        dz.pack(fill="x", padx=12, pady=(0, 6))
        ctk.CTkLabel(dz, text="📂  Click Browse or drag & drop",
                     font=FONT_SMALL, text_color=ACCENT,
                     anchor="center").pack(pady=(10, 2))
        tag_row = ctk.CTkFrame(dz, fg_color="transparent")
        tag_row.pack(pady=(0, 8))
        for tag in [".json", ".smap"]:
            ctk.CTkLabel(tag_row, text=tag, font=FONT_SMALL, text_color=TEXT_LO,
                         fg_color=BG_PANEL, corner_radius=4,
                         width=44).pack(side="left", padx=3)
        fe_row = ctk.CTkFrame(c1, fg_color="transparent")
        fe_row.pack(fill="x", padx=12, pady=(0, 4))
        self._tab2_file_entry = ctk.CTkEntry(
            fe_row, placeholder_text="No file selected",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            height=32, corner_radius=8, font=FONT_SMALL)
        self._tab2_file_entry.pack(side="left", fill="x", expand=True, padx=(0, 6))
        ctk.CTkButton(fe_row, text="Browse", width=70,
                      fg_color=BG_PANEL, hover_color=BORDER, border_color=BORDER,
                      border_width=1, text_color=TEXT_MID, corner_radius=8,
                      height=32, font=FONT_SMALL,
                      command=self._tab2_browse_file).pack(side="left")
        GhostButton(c1, text="⟳  Reload Map",
                    command=lambda: self._map_load_from_path(
                        self._tab2_file_entry.get())).pack(
            fill="x", padx=12, pady=(0, 12))

        # ── Card 2: Orientation + Detection ──
        ctk.CTkLabel(c2, text="ORIENTATION", font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(
            anchor="w", padx=14, pady=(12, 6))
        self.orientation_var = tk.StringVar(value="horizontal")
        orient_seg = ctk.CTkSegmentedButton(
            c2, values=["Horizontal", "Vertical"],
            command=lambda v: self.orientation_var.set(v.lower()),
            fg_color=BG_INPUT, selected_color=ACCENT,
            selected_hover_color=ACCENT_ALT, unselected_color=BG_INPUT,
            unselected_hover_color=BORDER, text_color=TEXT_MID,
            font=FONT_SMALL, height=32, corner_radius=7)
        orient_seg.set("Horizontal")
        orient_seg.pack(fill="x", padx=12, pady=(0, 10))
        ctk.CTkLabel(c2, text="DETECTION", font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(
            anchor="w", padx=14, pady=(4, 6))
        self.cross_aisle_var = ctk.BooleanVar(value=True)
        ctk.CTkCheckBox(c2, text="Enable cross-aisle detection",
                        variable=self.cross_aisle_var,
                        font=FONT_SMALL, text_color=TEXT_MID,
                        fg_color=ACCENT, hover_color=ACCENT_ALT,
                        checkmark_color="#fff",
                        border_color=BORDER).pack(
            anchor="w", padx=14, pady=(0, 14))

        # Hidden state vars — kept for sequence routing logic, not shown in UI
        self.mode_var = tk.StringVar(value="automatic")
        self.loading_point_var = tk.StringVar()
        self.loading_combo = ctk.CTkComboBox(
            p, variable=self.loading_point_var, values=[], state="readonly",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            button_color=BORDER, button_hover_color=ACCENT,
            dropdown_fg_color=BG_PANEL, dropdown_text_color=TEXT_MID,
            font=FONT_LABEL, height=34, corner_radius=8)
        self.tab2_ap_count_label = ctk.CTkLabel(p, text="—")
        # intentionally not packed — both widgets are hidden from UI

        # ── Card 4: Output & Duplication ──
        ctk.CTkLabel(c4, text="OUTPUT & DUPLICATION",
                     font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(
            anchor="w", padx=14, pady=(12, 8))
        tmpl_row = ctk.CTkFrame(c4, fg_color="transparent")
        tmpl_row.pack(fill="x", padx=12, pady=(0, 6))
        ctk.CTkLabel(tmpl_row, text="Template:", font=FONT_SMALL,
                     text_color=TEXT_MID, width=66, anchor="w").pack(side="left")
        self._tab2_template_entry = ctk.CTkEntry(
            tmpl_row, placeholder_text="ra_input.xlsx",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            height=30, corner_radius=8, font=FONT_SMALL)
        self._tab2_template_entry.pack(side="left", fill="x", expand=True, padx=(0, 4))
        ctk.CTkButton(tmpl_row, text="…", width=28,
                      fg_color=BG_PANEL, hover_color=BORDER, border_color=BORDER,
                      border_width=1, text_color=TEXT_MID, corner_radius=8,
                      height=30, font=FONT_SMALL,
                      command=self._tab2_browse_template).pack(side="left")
        out_row = ctk.CTkFrame(c4, fg_color="transparent")
        out_row.pack(fill="x", padx=12, pady=(0, 8))
        ctk.CTkLabel(out_row, text="Output:", font=FONT_SMALL,
                     text_color=TEXT_MID, width=66, anchor="w").pack(side="left")
        self.tab2_output_entry = ctk.CTkEntry(
            out_row, placeholder_text="pick_sequences.xlsx",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            height=30, corner_radius=8, font=FONT_SMALL)
        self.tab2_output_entry.insert(0, "pick_sequences.xlsx")
        self.tab2_output_entry.pack(side="left", fill="x", expand=True)
        tk.Frame(c4, bg=BORDER, height=1).pack(fill="x", padx=12, pady=(4, 8))
        self.dup_enabled_var = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(c4, text="Enable duplication",
                        variable=self.dup_enabled_var,
                        command=self._toggle_dup_options,
                        font=FONT_SMALL, text_color=TEXT_MID,
                        fg_color=ACCENT, hover_color=ACCENT_ALT,
                        checkmark_color="#fff",
                        border_color=BORDER).pack(
            anchor="w", padx=14, pady=(0, 6))
        self._dup_options_frame = ctk.CTkFrame(c4, fg_color="transparent")
        dup_inner = ctk.CTkFrame(self._dup_options_frame, fg_color="transparent")
        dup_inner.pack(fill="x", padx=14)
        ctk.CTkLabel(dup_inner, text="Count:", font=FONT_SMALL,
                     text_color=TEXT_MID, width=50, anchor="w").pack(side="left")
        self.dup_count_entry = ctk.CTkEntry(
            dup_inner, width=70, fg_color=BG_INPUT, border_color=BORDER,
            text_color=TEXT_HI, height=30, corner_radius=8, font=FONT_SMALL)
        self.dup_count_entry.insert(0, "2")
        self.dup_count_entry.pack(side="left", padx=(6, 0))

        # Action buttons
        act_row = ctk.CTkFrame(sec1, fg_color="transparent")
        act_row.pack(fill="x", padx=20, pady=(12, 16))
        GhostButton(act_row, text="🔍  Load & Analyze",
                    command=self.load_and_analyze).pack(side="left", padx=(0, 8))
        AccentButton(act_row, text="🚀  Generate Sequence",
                     command=self.generate_sequence).pack(side="left")

        # ══ SECTION 2 — Warehouse Map & Aisle Directions ══
        sec2 = ctk.CTkFrame(outer, fg_color=BG_PANEL, corner_radius=12,
                             border_color=BORDER, border_width=1)
        sec2.pack(fill="x", padx=20, pady=(16, 16))

        # Section 2 header
        s2_hdr = ctk.CTkFrame(sec2, fg_color="transparent")
        s2_hdr.pack(fill="x", padx=20, pady=(14, 0))
        ctk.CTkLabel(s2_hdr, text="2", width=22, height=22,
                     fg_color=ACCENT, text_color="#fff", corner_radius=11,
                     font=("Segoe UI", 10, "bold")).pack(side="left", padx=(0, 10))
        s2_t = ctk.CTkFrame(s2_hdr, fg_color="transparent")
        s2_t.pack(side="left")
        ctk.CTkLabel(s2_t, text="Warehouse Map & Aisle Directions",
                     font=("Segoe UI", 13, "bold"), text_color=TEXT_HI,
                     anchor="w").pack(anchor="w")
        ctk.CTkLabel(s2_t, text="Visual map preview and per-aisle direction control",
                     font=FONT_SMALL, text_color=TEXT_LO, anchor="w").pack(anchor="w")
        GhostButton(s2_hdr, text="⟳ Reset",
                    command=self._map_reset_view,
                    width=76).pack(side="right", padx=(4, 0))
        GhostButton(s2_hdr, text="✕ Clear",
                    command=self._map_clear_exclusions,
                    width=76).pack(side="right", padx=(4, 0))
        def _badge_frm(parent, prefix, color):
            f = ctk.CTkFrame(parent, fg_color=BG_INPUT, corner_radius=6)
            f.pack(side="right", padx=(0, 6))
            ctk.CTkLabel(f, text=f" {prefix}: ", font=FONT_SMALL,
                         text_color=TEXT_MID).pack(side="left")
            lbl = ctk.CTkLabel(f, text="0 ", font=("Segoe UI", 10, "bold"),
                               text_color=color)
            lbl.pack(side="left")
            return lbl
        self._map_excl_count_label   = _badge_frm(s2_hdr, "Excl",     ERROR)
        self._map_unload_count_label = _badge_frm(s2_hdr, "Unload",  WARNING)
        self._map_load_count_label   = _badge_frm(s2_hdr, "Load",    SUCCESS)
        self._map_ca_count_label     = _badge_frm(s2_hdr, "CA",      "#06B6D4")

        # ── Map box ──
        map_box = Card(sec2)
        map_box.pack(fill="x", padx=16, pady=(12, 0))
        map_toolbar = ctk.CTkFrame(map_box, fg_color="transparent")
        map_toolbar.pack(fill="x", padx=12, pady=(8, 4))
        for hint in ["Scroll = zoom", "Right drag = pan", "Left click/drag = tag"]:
            ctk.CTkLabel(map_toolbar, text=hint, font=FONT_SMALL,
                         text_color=TEXT_LO, fg_color=BG_INPUT,
                         corner_radius=12).pack(side="left", padx=(0, 6), pady=2)

        # Tag-mode selector
        tag_row = ctk.CTkFrame(map_box, fg_color="transparent")
        tag_row.pack(fill="x", padx=12, pady=(0, 6))
        ctk.CTkLabel(tag_row, text="Tag as:", font=FONT_SMALL,
                     text_color=TEXT_MID).pack(side="left", padx=(0, 8))
        self._map_tag_mode = tk.StringVar(value="remove")
        self._map_tag_btns = {}
        for val, lbl, color in [
            ("loading",     "🟢 Loading",    SUCCESS),
            ("unloading",   "🟠 Unloading",  WARNING),
            ("cross_aisle", "✚ Cross-Aisle", "#06B6D4"),
            ("remove",      "✕ Remove",      ERROR),
        ]:
            b = ctk.CTkButton(
                tag_row, text=lbl, width=100,
                fg_color=color if val == "remove" else BG_PANEL,
                hover_color=color, border_color=color, border_width=1,
                text_color=TEXT_HI if val == "remove" else color,
                corner_radius=7, height=28, font=FONT_SMALL,
                command=lambda v=val: self._map_set_tag_mode(v))
            b.pack(side="left", padx=(0, 5))
            self._map_tag_btns[val] = b
        self._map_show_ap     = ctk.BooleanVar(value=True)
        self._map_show_lm     = ctk.BooleanVar(value=True)
        self._map_show_labels = ctk.BooleanVar(value=True)
        for var, lbl, color in [
            (self._map_show_labels, "Labels", TEXT_LO),
            (self._map_show_lm,     "LMs",    SUCCESS),
            (self._map_show_ap,     "APs",    ACCENT),
        ]:
            ctk.CTkCheckBox(map_toolbar, text=lbl, variable=var,
                            command=self._map_redraw,
                            font=FONT_SMALL, text_color=color,
                            fg_color=color, hover_color=color,
                            checkmark_color="#fff", border_color=BORDER,
                            checkbox_width=16, checkbox_height=16).pack(
                side="right", padx=(10, 0))
        canvas_outer = tk.Frame(map_box, bg=BG_INPUT, height=300)
        canvas_outer.pack(fill="x", padx=12, pady=(0, 10))
        canvas_outer.pack_propagate(False)
        self._map_canvas = tk.Canvas(
            canvas_outer, bg=BG_INPUT, highlightthickness=0, cursor="crosshair")
        self._map_canvas.pack(fill="both", expand=True)
        self._map_canvas.bind("<ButtonPress-1>",   self._map_press)
        self._map_canvas.bind("<B1-Motion>",       self._map_drag_select)
        self._map_canvas.bind("<ButtonRelease-1>", self._map_release)
        self._map_canvas.bind("<ButtonPress-3>",   self._map_pan_start)
        self._map_canvas.bind("<B3-Motion>",       self._map_pan_move)
        self._map_canvas.bind("<ButtonRelease-3>", self._map_pan_end)
        self._map_canvas.bind("<MouseWheel>",      self._map_zoom)
        self._map_canvas.bind("<Configure>",       self._map_on_resize)
        # Prevent wheel events from reaching the outer scrollable frame
        self._map_canvas.bind("<Enter>", lambda _: self._map_canvas.bind_all(
            "<MouseWheel>", self._map_zoom))
        self._map_canvas.bind("<Leave>", lambda _: self._map_canvas.unbind_all(
            "<MouseWheel>"))

        # ── Bottom row: Aisle Directions (left) + Log (right) ──
        bot_row = tk.Frame(sec2, bg=BG_PANEL)
        bot_row.pack(fill="x", padx=16, pady=(12, 16))
        bot_row.columnconfigure(0, weight=1)
        bot_row.columnconfigure(1, weight=1)

        aisle_card = Card(bot_row)
        aisle_card.grid(row=0, column=0, padx=(0, 8), sticky="nsew")
        aisle_hdr = ctk.CTkFrame(aisle_card, fg_color="transparent")
        aisle_hdr.pack(fill="x", padx=14, pady=(12, 4))
        ctk.CTkLabel(aisle_hdr, text="Aisle Directions",
                     font=FONT_HEAD, text_color=TEXT_HI).pack(side="left")
        for txt, cmd in [
            ("All F→L", lambda: self.set_all_directions("f2l")),
            ("All L→F", lambda: self.set_all_directions("l2f")),
            ("Toggle",  self.toggle_selected_directions),
        ]:
            GhostButton(aisle_hdr, text=txt, command=cmd,
                        width=72).pack(side="right", padx=(4, 0))
        tree_outer = ctk.CTkFrame(aisle_card, fg_color="transparent")
        tree_outer.pack(fill="x", padx=10, pady=(4, 10))
        tree_wrap, self.aisle_tree = make_styled_tree(
            tree_outer,
            columns=("Aisle", "First AP", "Last AP", "Count", "Direction"),
            headings=("Aisle #", "First AP", "Last AP", "APs", "Direction"),
            col_widths=(70, 160, 160, 50, 110))
        tree_wrap.pack(fill="x")

        log_card = Card(bot_row)
        log_card.grid(row=0, column=1, sticky="nsew")
        log_hdr = ctk.CTkFrame(log_card, fg_color="transparent")
        log_hdr.pack(fill="x", padx=14, pady=(12, 4))
        ctk.CTkLabel(log_hdr, text="Log",
                     font=FONT_HEAD, text_color=TEXT_HI).pack(side="left")
        ctk.CTkLabel(log_hdr, text="●", font=("Segoe UI", 12),
                     text_color=SUCCESS).pack(side="right")
        self.tab2_log = LogBox(log_card, height=160)
        self.tab2_log.pack(fill="x", padx=12, pady=(0, 10))

        # Map state
        self._map_points          = []
        self._map_excluded        = set()
        self._map_loading_aps     = set()
        self._map_unloading_aps   = set()
        self._map_cross_aisle_lms = set()
        self._map_scale         = 1.0
        self._map_pan_x         = 0.0
        self._map_pan_y         = 0.0
        self._map_world_cx      = 0.0
        self._map_world_cy      = 0.0
        self._map_drag_last     = None
        self._map_fitted        = False
        self._map_sel_start     = None
        self._map_sel_active    = False

    def log_tab2(self, msg: str):
        tag = "ok" if "✅" in msg else "err" if "❌" in msg else \
              "warn" if "⚠️" in msg else "info"
        clean = msg.lstrip("✅❌⚠️💥🔄🔍🔗📊👤🤖🚫").strip()
        self.tab2_log.log(clean, tag)

    # ─────────────────── Tab 3: Zone Configuration ───────────────────

    def setup_tab3(self):
        p = self.pages["tab3"]
        outer = ctk.CTkScrollableFrame(p, fg_color="transparent",
                                        scrollbar_button_color=BORDER,
                                        scrollbar_button_hover_color=ACCENT)
        outer.pack(fill="both", expand=True, padx=24, pady=20)

        ctk.CTkLabel(outer, text="Point to Point",
                     font=FONT_TITLE, text_color=TEXT_HI,
                     anchor="w").pack(anchor="w", pady=(0, 18))

        # Input files card
        c1 = Card(outer, title="Input Files")
        c1.pack(fill="x", pady=(0, 14))
        f1 = ctk.CTkFrame(c1, fg_color="transparent")
        f1.pack(fill="x", padx=16, pady=(10, 16))

        self._tab3_json_picker = FilePickerRow(
            f1, "Mapping JSON",
            [("JSON / SMAP", "*.json *.smap"), ("JSON", "*.json"), ("SMAP", "*.smap")],
            on_browse=self._on_tab3_json_browse)
        self._tab3_json_picker.pack(fill="x", pady=(0, 8))

        self._tab3_excel_picker = FilePickerRow(
            f1, "Excel Config", [("Excel", "*.xlsx")],
            on_browse=self._on_tab3_excel_browse)
        self._tab3_excel_picker.pack(fill="x", pady=(0, 8))

        sample_row = ctk.CTkFrame(f1, fg_color="transparent")
        sample_row.pack(anchor="w")
        ctk.CTkLabel(sample_row, text="Not sure about the format?",
                     font=FONT_SMALL, text_color=TEXT_LO).pack(side="left", padx=(0, 8))
        GhostButton(sample_row, text="⬇  Download Sample Excel",
                    command=self._download_sample_excel,
                    width=200).pack(side="left")

        # Entry/exit config card
        c2 = Card(outer, title="Entry / Exit Points")
        c2.pack(fill="x", pady=(0, 14))
        f2 = ctk.CTkFrame(c2, fg_color="transparent")
        f2.pack(fill="x", padx=16, pady=(10, 16))
        ctk.CTkLabel(f2, text="Location Entry/Exit Point Type",
                     font=FONT_LABEL, text_color=TEXT_MID).pack(
            anchor="w", pady=(0, 6))
        entry_row = ctk.CTkFrame(f2, fg_color="transparent")
        entry_row.pack(anchor="w")
        self.location_entry_type = tk.StringVar(value="LM")
        for val, lbl in [("LM", "Use Location Marks (LM)"),
                          ("AP", "Use Action Points (AP)")]:
            ctk.CTkRadioButton(entry_row, text=lbl, value=val,
                               variable=self.location_entry_type,
                               font=FONT_LABEL, text_color=TEXT_MID,
                               fg_color=ACCENT, hover_color=ACCENT_ALT,
                               border_color=BORDER).pack(
                side="left", padx=(0, 20))

        # ── N-Deep Configuration card ──────────────────────────────
        c_nd = Card(outer, title="N-Deep Configuration")
        c_nd.pack(fill="x", pady=(0, 14))
        f_nd = ctk.CTkFrame(c_nd, fg_color="transparent")
        f_nd.pack(fill="x", padx=16, pady=(10, 16))

        self.ndeep_enabled_var = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(
            f_nd,
            text="Enable N-Deep zone generation (auto-generates subzones from JSON path connectivity)",
            variable=self.ndeep_enabled_var,
            command=self._on_ndeep_toggle,
            font=FONT_LABEL, text_color=TEXT_MID,
            fg_color=ACCENT, hover_color=ACCENT_ALT,
            checkmark_color="#fff", border_color=BORDER
        ).pack(anchor="w", pady=(0, 10))

        # Hidden options frame — revealed when checkbox is ticked
        self._ndeep_options_frame = ctk.CTkFrame(f_nd, fg_color="transparent")

        # Drop sequence
        ctk.CTkLabel(self._ndeep_options_frame, text="Drop Sequence",
                     font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(anchor="w", pady=(0, 4))
        self.ndeep_drop_seq_var = tk.StringVar(value="f2l")
        seq_row = ctk.CTkFrame(self._ndeep_options_frame, fg_color="transparent")
        seq_row.pack(anchor="w", pady=(0, 10))
        for val, lbl in [
            ("f2l", "First to Last  (AP nearest entry = index 001)"),
            ("l2f", "Last to First  (AP farthest from entry = index 001)"),
        ]:
            ctk.CTkRadioButton(
                seq_row, text=lbl, value=val,
                variable=self.ndeep_drop_seq_var,
                font=FONT_LABEL, text_color=TEXT_MID,
                fg_color=ACCENT, hover_color=ACCENT_ALT,
                border_color=BORDER
            ).pack(side="left", padx=(0, 20))

        # Entry / Exit scope — three independent levels
        ctk.CTkLabel(self._ndeep_options_frame, text="Entry / Exit Scope",
                     font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(anchor="w", pady=(4, 2))

        def _scope_row(parent, label, var):
            row_frame = ctk.CTkFrame(parent, fg_color="transparent")
            row_frame.pack(anchor="w", fill="x", pady=(2, 4))
            ctk.CTkLabel(row_frame, text=label, font=FONT_LABEL,
                         text_color=TEXT_MID, width=90, anchor="w").pack(side="left")
            for val, lbl in [("BOTH", "Both"), ("ENTRY", "Entry Only"), ("EXIT", "Exit Only")]:
                ctk.CTkRadioButton(
                    row_frame, text=lbl, value=val, variable=var,
                    font=FONT_LABEL, text_color=TEXT_MID,
                    fg_color=ACCENT, hover_color=ACCENT_ALT,
                    border_color=BORDER
                ).pack(side="left", padx=(0, 16))

        self.ndeep_zone_scope_var = tk.StringVar(value="BOTH")
        self.ndeep_sub_scope_var  = tk.StringVar(value="BOTH")
        self.ndeep_loc_scope_var  = tk.StringVar(value="BOTH")
        _scope_row(self._ndeep_options_frame, "Zone:",     self.ndeep_zone_scope_var)
        _scope_row(self._ndeep_options_frame, "Subzone:",  self.ndeep_sub_scope_var)
        _scope_row(self._ndeep_options_frame, "Location:", self.ndeep_loc_scope_var)

        ctk.CTkLabel(
            self._ndeep_options_frame,
            text="Excel columns needed: Zone name | zone entry | zone exit | min x | max x | min y | max y\n"
                 ".smap files are auto-converted to JSON on selection.",
            font=FONT_SMALL, text_color=TEXT_LO, justify="left"
        ).pack(anchor="w", pady=(0, 4))

        # Sequence config card
        c3 = Card(outer, title="Sequence Configuration")
        c3.pack(fill="x", pady=(0, 14))
        f3 = ctk.CTkFrame(c3, fg_color="transparent")
        f3.pack(fill="x", padx=16, pady=(10, 16))

        self.create_sequence_var = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(f3, text="Create pick sequences for subzones",
                        variable=self.create_sequence_var,
                        command=self.toggle_sequence_options,
                        font=FONT_LABEL, text_color=TEXT_MID,
                        fg_color=ACCENT, hover_color=ACCENT_ALT,
                        checkmark_color="#fff",
                        border_color=BORDER).pack(anchor="w", pady=(0, 10))

        # Subzone treeview (hidden until checkbox ticked)
        self.sequence_options_frame = ctk.CTkFrame(f3, fg_color="transparent")

        ctk.CTkLabel(self.sequence_options_frame,
                     text="Configure subzone directions:",
                     font=FONT_LABEL, text_color=TEXT_MID,
                     anchor="w").pack(anchor="w", pady=(0, 6))

        sub_tree_wrap, self.subzone_tree = make_styled_tree(
            self.sequence_options_frame,
            columns=("Subzone", "Direction"),
            headings=("Subzone Name", "Direction"),
            col_widths=(280, 200))
        sub_tree_wrap.pack(fill="x", ipady=80)

        sub_dir_row = ctk.CTkFrame(self.sequence_options_frame,
                                   fg_color="transparent")
        sub_dir_row.pack(anchor="w", pady=(8, 0))
        for txt, cmd_arg in [
            ("← Left→Right",  "left_to_right"),
            ("→ Right→Left",  "right_to_left"),
            ("↑ Top→Bottom",  "top_to_bottom"),
            ("↓ Bottom→Top",  "bottom_to_top"),
        ]:
            GhostButton(sub_dir_row, text=txt, width=120,
                        command=lambda a=cmd_arg: self.set_all_subzone_directions(a)
                        ).pack(side="left", padx=(0, 6))
        GhostButton(sub_dir_row, text="🔁 Toggle Selected",
                    command=self.toggle_subzone_directions,
                    width=140).pack(side="left")
        self.subzone_tree.bind("<Double-1>", self.on_subzone_direction_click)

        # Output
        c4 = Card(outer, title="Output File")
        c4.pack(fill="x", pady=(0, 14))
        f4 = ctk.CTkFrame(c4, fg_color="transparent")
        f4.pack(fill="x", padx=16, pady=(10, 16))
        self.tab3_output_entry = ctk.CTkEntry(
            f4, placeholder_text="zone_config.json",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            height=34, corner_radius=8, font=FONT_LABEL)
        self.tab3_output_entry.insert(0, "zone_config.json")
        self.tab3_output_entry.pack(fill="x")

        AccentButton(outer, text="🚀  Generate Zone Configuration",
                     command=self.run_zone_configuration,
                     width=260).pack(anchor="w", pady=(8, 18))

        log_card = Card(outer, title="Log")
        log_card.pack(fill="x")
        self.tab3_log = LogBox(log_card, height=200)
        self.tab3_log.pack(fill="x", padx=16, pady=(8, 16))

    def _on_tab3_excel_browse(self, path):
        self.show_excel_format_info()
        if self.create_sequence_var.get():
            self.load_subzones_from_excel()

    def _download_sample_excel(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile="point_to_point_sample.xlsx",
            filetypes=[("Excel", "*.xlsx")],
            title="Save Sample Excel")
        if not path:
            return
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook()

            # ── Sheet 1: Standard mode (with subzones) ──────────────────────
            ws1 = wb.active
            ws1.title = "Standard"

            std_headers = [
                "Zone name", "zone entry", "zone exit",
                "subzone name", "subzone entery", "subzone exit",
                "min x", "max x", "min y", "max y",
            ]
            std_rows = [
                ["Zone_A", "LM_A_Entry", "LM_A_Exit", "Sub_A1", "LM_Sub_A1_Entry", "LM_Sub_A1_Exit",  0.0, 10.0,  0.0,  5.0],
                ["Zone_A", "LM_A_Entry", "LM_A_Exit", "Sub_A2", "LM_Sub_A2_Entry", "LM_Sub_A2_Exit",  0.0, 10.0,  5.0, 10.0],
                ["Zone_A", "LM_A_Entry", "LM_A_Exit", "Sub_A3", "LM_Sub_A3_Entry", "LM_Sub_A3_Exit",  0.0, 10.0, 10.0, 15.0],
                ["Zone_B", "LM_B_Entry", "LM_B_Exit", "Sub_B1", "LM_Sub_B1_Entry", "LM_Sub_B1_Exit", 10.0, 20.0,  0.0,  5.0],
                ["Zone_B", "LM_B_Entry", "LM_B_Exit", "Sub_B2", "LM_Sub_B2_Entry", "LM_Sub_B2_Exit", 10.0, 20.0,  5.0, 10.0],
            ]

            hdr_font = Font(bold=True, color="FFFFFF")
            hdr_fill = PatternFill("solid", fgColor="2563EB")
            alt_fill = PatternFill("solid", fgColor="1E293B")

            ws1.append(std_headers)
            for cell in ws1[1]:
                cell.font  = hdr_font
                cell.fill  = hdr_fill
                cell.alignment = Alignment(horizontal="center")

            for i, row in enumerate(std_rows):
                ws1.append(row)
                if i % 2 == 1:
                    for cell in ws1[i + 2]:
                        cell.fill = alt_fill

            for col in ws1.columns:
                ws1.column_dimensions[col[0].column_letter].width = max(
                    len(str(col[0].value or "")), 14) + 2

            # ── Sheet 2: N-Deep mode (no subzone columns) ───────────────────
            ws2 = wb.create_sheet("N-Deep")

            nd_headers = [
                "Zone name", "zone entry", "zone exit",
                "min x", "max x", "min y", "max y",
            ]
            nd_rows = [
                ["Zone_A", "LM_A_Entry", "LM_A_Exit",  0.0, 10.0,  0.0, 15.0],
                ["Zone_B", "LM_B_Entry", "LM_B_Exit", 10.0, 20.0,  0.0, 10.0],
                ["Zone_C", "LM_C_Entry", "LM_C_Exit", 20.0, 30.0,  0.0, 10.0],
            ]

            ws2.append(nd_headers)
            for cell in ws2[1]:
                cell.font  = hdr_font
                cell.fill  = hdr_fill
                cell.alignment = Alignment(horizontal="center")

            for i, row in enumerate(nd_rows):
                ws2.append(row)
                if i % 2 == 1:
                    for cell in ws2[i + 2]:
                        cell.fill = alt_fill

            for col in ws2.columns:
                ws2.column_dimensions[col[0].column_letter].width = max(
                    len(str(col[0].value or "")), 14) + 2

            wb.save(path)
            messagebox.showinfo(
                "Sample Downloaded",
                f"Sample saved to:\n{path}\n\n"
                "Sheet 'Standard'  — use when N-Deep is OFF\n"
                "Sheet 'N-Deep'    — use when N-Deep is ON\n\n"
                "Replace the marker names with real instanceNames from your .json/.smap file.\n"
                "Coordinates (min x / max x / min y / max y) define the bounding box of each zone/subzone.")
        except Exception as e:
            messagebox.showerror("Error", f"Could not create sample file:\n{e}")

    # ── N-deep toggle / JSON browse / SMAP conversion ───────────────

    def _on_ndeep_toggle(self):
        if self.ndeep_enabled_var.get():
            self._ndeep_options_frame.pack(fill="x", pady=(0, 8))
        else:
            self._ndeep_options_frame.pack_forget()

    def _on_tab3_json_browse(self, path: str):
        if path.lower().endswith(".smap"):
            self.log_tab3(f"⚙️ .smap detected — converting to JSON…")
            converted = self._convert_smap_to_json(path)
            if converted:
                self._tab3_json_picker.set(converted)
                self.log_tab3(f"✅ Converted → {os.path.basename(converted)}")
            else:
                self.log_tab3("❌ SMAP conversion failed")
        else:
            self.show_json_format_info()

    def _convert_smap_to_json(self, smap_path: str):
        """Load a .smap file (assumed JSON) and write a _converted.json alongside it."""
        try:
            with open(smap_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            base = os.path.splitext(smap_path)[0]
            out_path = base + "_converted.json"
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2)
            return out_path
        except json.JSONDecodeError as e:
            messagebox.showerror("SMAP Parse Error",
                                 f"Cannot parse .smap file as JSON:\n{e}")
            return None
        except Exception as e:
            messagebox.showerror("SMAP Error", f"Cannot convert .smap file:\n{e}")
            return None

    # ── N-deep Excel validation ──────────────────────────────────────

    def validate_excel_for_ndeep(self, df):
        required = ['Zone name', 'zone entry', 'zone exit',
                    'min x', 'max x', 'min y', 'max y']
        missing = [c for c in required if c not in df.columns]
        if missing:
            return False, f"Missing columns for N-Deep: {missing}"
        try:
            for col in ['min x', 'max x', 'min y', 'max y']:
                pd.to_numeric(df[col], errors='raise')
        except Exception as e:
            return False, f"Coordinate columns must be numeric: {e}"
        return True, "OK"

    # ── N-deep subzone grouping ──────────────────────────────────────

    def _group_aps_into_inline_subzones(self, zone_aps, graph, node_positions,
                                        tolerance=0.3):
        """
        Group APs in a zone into inline subzones (one subzone = one rack lane).

        Algorithm:
          1. Try grouping by same-Y (horizontal lanes) and by same-X (vertical lanes).
          2. Score each orientation by counting directly-connected consecutive AP pairs
             (direct graph edges) — the orientation with more direct connections wins.
          3. Within each spatial group, further split wherever two consecutive APs are
             NOT reachable from each other without passing through another zone AP
             (i.e. there is a physical break in the lane).
        """
        if not zone_aps:
            return []
        if len(zone_aps) == 1:
            return [zone_aps]

        # Step 1 — try both orientations
        y_groups = self._spatial_group(zone_aps, primary='y', secondary='x',
                                       tolerance=tolerance)
        x_groups = self._spatial_group(zone_aps, primary='x', secondary='y',
                                       tolerance=tolerance)

        # Step 2 — pick the orientation whose consecutive pairs are more often
        #           directly connected in the graph (highest score wins)
        y_score = self._score_grouping(y_groups, graph)
        x_score = self._score_grouping(x_groups, graph)
        raw_groups = y_groups if y_score >= x_score else x_groups

        # Step 3 — within each spatial group split on path breaks
        zone_ap_names = {ap["name"] for ap in zone_aps}
        final_groups = []
        for group in raw_groups:
            subgroups = self._split_consecutive(group, graph, zone_ap_names)
            final_groups.extend(subgroups)

        return [g for g in final_groups if g]

    def _spatial_group(self, aps, primary='y', secondary='x', tolerance=0.3):
        """Group APs with approximately the same primary-axis value; sort by secondary."""
        sorted_aps = sorted(aps, key=lambda a: (a[primary], a[secondary]))
        groups, current = [], [sorted_aps[0]]
        ref = sorted_aps[0][primary]
        for ap in sorted_aps[1:]:
            if abs(ap[primary] - ref) <= tolerance:
                current.append(ap)
            else:
                groups.append(sorted(current, key=lambda a: a[secondary]))
                current = [ap]
                ref = ap[primary]
        groups.append(sorted(current, key=lambda a: a[secondary]))
        return groups

    def _score_grouping(self, groups, graph):
        """Count direct graph-edge connections between consecutive APs in each group."""
        score = 0
        for group in groups:
            for i in range(len(group) - 1):
                a, b = group[i]["name"], group[i + 1]["name"]
                if b in graph.get(a, {}):
                    score += 1
        return score

    def _split_consecutive(self, sorted_aps, graph, zone_ap_names, factor=2.5):
        """
        Walk the sorted AP list and start a new subgroup wherever the lane breaks.

        A break is detected when ANY of the following are true for consecutive (i, i+1):
          1. The Euclidean gap is anomalously large compared to the minimum gap
             in this spatial group  (gap > 2.5 × min_gap of the group).
          2. There is no direct graph edge AND the restricted path distance
             (path through LMs only, not other zone APs) exceeds factor × Euclidean.
          3. There is no direct graph edge AND no restricted path exists at all.

        Priority order: direct edge (connect) → anomalous gap (split) → path check.
        """
        if len(sorted_aps) <= 1:
            return [sorted_aps]

        # Pre-compute all consecutive Euclidean distances for gap-anomaly detection
        eucl_gaps = [
            math.hypot(sorted_aps[i + 1]["x"] - sorted_aps[i]["x"],
                       sorted_aps[i + 1]["y"] - sorted_aps[i]["y"])
            for i in range(len(sorted_aps) - 1)
        ]
        min_gap = min(eucl_gaps)
        # Threshold: a gap is anomalous if it is more than 2.5× the smallest gap
        gap_threshold = max(min_gap * 2.5, 0.5)

        subgroups, current = [], [sorted_aps[0]]

        for i, eucl in enumerate(eucl_gaps):
            prev_ap = sorted_aps[i]
            curr_ap = sorted_aps[i + 1]
            prev_name = prev_ap["name"]
            curr_name = curr_ap["name"]

            # ── Check 1: direct graph edge ───────────────────────────
            if curr_name in graph.get(prev_name, {}):
                current.append(curr_ap)
                continue

            # ── Check 2: spatial gap anomaly → force split ───────────
            if eucl > gap_threshold:
                subgroups.append(current)
                current = [curr_ap]
                continue

            # ── Check 3: short path through LMs only ─────────────────
            excluded = zone_ap_names - {prev_name, curr_name}
            restricted = {
                n: {v: w for v, w in nb.items() if v not in excluded}
                for n, nb in graph.items()
                if n not in excluded
            }
            connected = False
            if prev_name in restricted:
                dists, _ = self.dijkstra(prev_name, restricted)
                path_d = dists.get(curr_name, float('inf'))
                if (not math.isinf(path_d)
                        and path_d <= factor * max(eucl, 0.01)):
                    connected = True

            if connected:
                current.append(curr_ap)
            else:
                subgroups.append(current)
                current = [curr_ap]

        subgroups.append(current)
        return subgroups

    def _order_subzone_aps(self, subzone_aps, zone_entry_name,
                            graph, node_positions, drop_seq="f2l"):
        """
        Order APs in a subzone by path distance from zone_entry_name.
        f2l → nearest AP gets index 001 (first in list).
        l2f → farthest AP gets index 001 (list is reversed).
        """
        dists, _ = self.dijkstra(zone_entry_name, graph)

        def ap_key(ap):
            d = dists.get(ap["name"], float('inf'))
            if math.isinf(d):
                ref = node_positions.get(zone_entry_name)
                if ref:
                    d = math.hypot(ap["x"] - ref[0], ap["y"] - ref[1])
            return d

        ordered = sorted(subzone_aps, key=ap_key)
        if drop_seq == "l2f":
            ordered = list(reversed(ordered))
        return ordered

    def _find_nearest_lm_to_ap(self, ap_name, lms_raw, graph, node_positions):
        """Return the nearest LocationMark name to ap_name by graph path distance."""
        # lms_raw contains dicts with {"name", "x", "y", "_raw"}; convert to the
        # format expected by find_closest_point_by_path_distance (advancedPointList items)
        lm_points = [r["_raw"] for r in lms_raw]
        return self.find_closest_point_by_path_distance(
            ap_name, lm_points, graph, node_positions, "LM")

    # ── N-deep main logic ────────────────────────────────────────────

    def generate_ndeep_zone_configuration(self, mapping_data, df,
                                           drop_seq,
                                           zone_scope, sub_scope, loc_scope,
                                           entry_exit_type,
                                           graph, node_positions):
        """
        Generate the N-deep zone configuration JSON.

        drop_seq:        "f2l" (nearest AP → index 001) | "l2f" (farthest → index 001)
        zone_scope:      "BOTH" | "ENTRY" | "EXIT"  — controls zone entryPoint/exitPoint
        sub_scope:       "BOTH" | "ENTRY" | "EXIT"  — controls subzone entryPoint/exitPoint
        loc_scope:       "BOTH" | "ENTRY" | "EXIT"  — controls location entryPoint/exitPoint
        entry_exit_type: "LM" | "AP"                — how location entry/exit is calculated
        locationType is always "BOTH" (it describes the location capability, not the fields).
        """
        aps_raw, lms_raw = [], []
        for p in mapping_data.get("advancedPointList", []):
            cls = p.get("className")
            pos = p.get("pos", {})
            if "x" not in pos or "y" not in pos:
                continue
            rec = {
                "name": p.get("instanceName", ""),
                "x": float(pos["x"]),
                "y": float(pos["y"]),
                "_raw": p,
            }
            if cls == "ActionPoint":
                aps_raw.append(rec)
            elif cls == "LocationMark":
                lms_raw.append(rec)

        zones_out = []

        # Preserve Excel row order for zone names
        seen_zones = []
        for zone_name in df['Zone name']:
            if zone_name not in seen_zones:
                seen_zones.append(zone_name)

        for zone_name in seen_zones:
            zone_rows = df[df['Zone name'] == zone_name]
            row0 = zone_rows.iloc[0]
            zone_entry_name = str(row0['zone entry'])
            zone_exit_name  = str(row0['zone exit'])
            min_x = float(zone_rows['min x'].min())
            max_x = float(zone_rows['max x'].max())
            min_y = float(zone_rows['min y'].min())
            max_y = float(zone_rows['max y'].max())

            # Collect APs within zone bounding box
            zone_aps = [a for a in aps_raw
                        if min_x <= a["x"] <= max_x and min_y <= a["y"] <= max_y]
            if not zone_aps:
                self.log_tab3(f"⚠️ No APs found in zone '{zone_name}', skipping")
                continue

            self.log_tab3(f"🔍 Zone '{zone_name}': {len(zone_aps)} APs found")

            # Group APs into inline subzones via spatial + connectivity analysis
            subzone_groups = self._group_aps_into_inline_subzones(
                zone_aps, graph, node_positions)
            self.log_tab3(f"   → {len(subzone_groups)} subzone(s) detected")

            # Sort subzones: closest to zone entry gets index 1
            if zone_entry_name in graph:
                dists_entry, _ = self.dijkstra(zone_entry_name, graph)
                subzone_groups.sort(
                    key=lambda grp: min(
                        dists_entry.get(ap["name"], float('inf')) for ap in grp))

            subzones_out = []
            for sz_idx, sz_aps in enumerate(subzone_groups, 1):
                sz_name = f"{zone_name}_{sz_idx}"

                # Order APs within subzone by distance from zone entry
                ordered_aps = self._order_subzone_aps(
                    sz_aps, zone_entry_name, graph, node_positions, drop_seq)

                # Subzone boundary LMs (nearest LM to each end of the ordered lane)
                ap_entry_side = ordered_aps[0]
                ap_exit_side  = ordered_aps[-1]
                sz_entry_lm = self._find_nearest_lm_to_ap(
                    ap_entry_side["name"], lms_raw, graph, node_positions)
                sz_exit_lm  = self._find_nearest_lm_to_ap(
                    ap_exit_side["name"],  lms_raw, graph, node_positions)

                # Build location entries
                locations = []
                for loc_idx, ap in enumerate(ordered_aps, 1):
                    fleet_loc = ap["name"]
                    cust_loc  = f"{sz_name}-{loc_idx:03d}"

                    # Calculate entry/exit point values
                    if entry_exit_type == "LM":
                        ep = self._find_nearest_lm_to_ap(
                            fleet_loc, lms_raw, graph, node_positions)
                        xp = ep
                    else:
                        # AP mode: entryPoint = previous AP in sequence
                        # First AP has no predecessor → fall back to nearest LM
                        if loc_idx == 1:
                            ep = self._find_nearest_lm_to_ap(
                                fleet_loc, lms_raw, graph, node_positions)
                        else:
                            ep = ordered_aps[loc_idx - 2]["name"]
                        xp = ep

                    # locationType is always "BOTH" — it describes what the
                    # location supports, not which fields we emit
                    loc_dict = {
                        "customerLocation":     cust_loc,
                        "fleetManagerLocation": fleet_loc,
                        "locationType":         "BOTH",
                    }
                    if loc_scope in ("BOTH", "ENTRY"):
                        loc_dict["entryPoint"] = ep
                    if loc_scope in ("BOTH", "EXIT"):
                        loc_dict["exitPoint"] = xp

                    locations.append(loc_dict)

                sz_dict = {
                    "name": sz_name,
                    "locationSelectionStrategy": "SUBZONELIFO",
                    "locations": locations,
                }
                if sub_scope in ("BOTH", "ENTRY"):
                    sz_dict["entryPoint"] = sz_entry_lm
                if sub_scope in ("BOTH", "EXIT"):
                    sz_dict["exitPoint"] = sz_exit_lm

                subzones_out.append(sz_dict)

            zone_dict = {
                "name": zone_name,
                "locationSelectionStrategy": "SUBZONELIFO",
                "subZones": subzones_out,
            }
            if zone_scope in ("BOTH", "ENTRY"):
                zone_dict["entryPoint"] = zone_entry_name
            if zone_scope in ("BOTH", "EXIT"):
                zone_dict["exitPoint"] = zone_exit_name

            zones_out.append(zone_dict)

        return {"LocationDataMapping": {"zones": zones_out}}

    def log_tab3(self, msg: str):
        tag = "ok" if "✅" in msg else "err" if "❌" in msg else \
              "warn" if "⚠️" in msg else "info"
        clean = msg.lstrip("✅❌⚠️💥🔄🔍🔗💾📊📍⚙️").strip()
        self.tab3_log.log(clean, tag)

    # ─────────────────── Helper: compat shims ───────────────────
    # These let original logic access file paths via the new widgets.

    @property
    def tab1_file_entry(self):
        return _EntryShim(self._tab1_file_picker)

    @property
    def tab2_file_entry(self):
        return self._tab2_file_entry

    @property
    def tab3_json_entry(self):
        return _EntryShim(self._tab3_json_picker)

    @property
    def tab3_excel_entry(self):
        return _EntryShim(self._tab3_excel_picker)

    # ─────────────────── Tab 4: Map Validator ───────────────────

    def setup_tab4(self):
        p = self.pages["tab4"]
        outer = ctk.CTkScrollableFrame(p, fg_color="transparent",
                                        scrollbar_button_color=BORDER,
                                        scrollbar_button_hover_color=ACCENT)
        outer.pack(fill="both", expand=True)

        # Header
        hdr = ctk.CTkFrame(outer, fg_color=BG_PANEL, corner_radius=12,
                            border_color=BORDER, border_width=1)
        hdr.pack(fill="x", padx=20, pady=(16, 0))
        hdr_row = ctk.CTkFrame(hdr, fg_color="transparent")
        hdr_row.pack(fill="x", padx=20, pady=14)
        ctk.CTkLabel(hdr_row, text="Map Validator",
                     font=FONT_TITLE, text_color=TEXT_HI,
                     anchor="w").pack(side="left")
        ctk.CTkLabel(hdr_row,
                     text="Checks entry/exit paths and full reachability between all points",
                     font=FONT_SMALL, text_color=TEXT_LO,
                     anchor="e").pack(side="right")

        # File input card
        file_card = Card(outer)
        file_card.pack(fill="x", padx=20, pady=(12, 0))
        ctk.CTkLabel(file_card, text="MAP FILE", font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(
            anchor="w", padx=14, pady=(12, 6))
        fp_row = ctk.CTkFrame(file_card, fg_color="transparent")
        fp_row.pack(fill="x", padx=12, pady=(0, 12))
        self._tab4_file_entry = ctk.CTkEntry(
            fp_row, placeholder_text="Select .json or .smap file…",
            fg_color=BG_INPUT, border_color=BORDER, text_color=TEXT_HI,
            height=36, corner_radius=8, font=FONT_LABEL)
        self._tab4_file_entry.pack(side="left", fill="x", expand=True, padx=(0, 8))
        ctk.CTkButton(fp_row, text="Browse", width=90,
                      fg_color=BG_PANEL, hover_color=BORDER, border_color=BORDER,
                      border_width=1, text_color=TEXT_MID, corner_radius=8,
                      height=36, font=FONT_LABEL,
                      command=self._tab4_browse).pack(side="left", padx=(0, 8))
        AccentButton(fp_row, text="▶  Validate",
                     command=self._run_map_validation).pack(side="left")

        # Summary bar (hidden until run)
        self._tab4_summary = ctk.CTkFrame(outer, fg_color="transparent")
        self._tab4_summary.pack(fill="x", padx=20, pady=(10, 0))

        # Results tree card
        results_card = Card(outer)
        results_card.pack(fill="x", padx=20, pady=(12, 0))
        ctk.CTkLabel(results_card, text="ISSUES FOUND",
                     font=("Segoe UI", 10, "bold"),
                     text_color=TEXT_LO, anchor="w").pack(
            anchor="w", padx=14, pady=(12, 6))
        tree_outer = ctk.CTkFrame(results_card, fg_color="transparent")
        tree_outer.pack(fill="x", padx=12, pady=(0, 12))
        tree_wrap, self._tab4_tree = make_styled_tree(
            tree_outer,
            columns=("Point", "Type", "Issue"),
            headings=("Point Name", "Type", "Issue"),
            col_widths=(220, 80, 340))
        self._tab4_tree.tag_configure("warn", foreground=WARNING)
        self._tab4_tree.tag_configure("err",  foreground=ERROR)
        self._tab4_tree.tag_configure("conn", foreground="#5B9BD5")
        tree_wrap.pack(fill="x")

        # Log
        log_card = Card(outer)
        log_card.pack(fill="x", padx=20, pady=(12, 16))
        log_hdr = ctk.CTkFrame(log_card, fg_color="transparent")
        log_hdr.pack(fill="x", padx=14, pady=(12, 4))
        ctk.CTkLabel(log_hdr, text="Log", font=FONT_HEAD,
                     text_color=TEXT_HI).pack(side="left")
        ctk.CTkLabel(log_hdr, text="●", font=("Segoe UI", 12),
                     text_color=SUCCESS).pack(side="right")
        self._tab4_log = LogBox(log_card, height=120)
        self._tab4_log.pack(fill="x", padx=12, pady=(0, 10))

    def _tab4_browse(self):
        path = filedialog.askopenfilename(
            filetypes=[("Map files", "*.json *.smap"),
                       ("JSON", "*.json"), ("SMAP", "*.smap")])
        if path:
            self._tab4_file_entry.delete(0, "end")
            self._tab4_file_entry.insert(0, path)

    def _check_connectivity(self, all_points: dict, outgoing: dict):
        """
        Tarjan's iterative SCC on the map graph, then BFS reachability on the
        condensation DAG. Returns (issue_rows, n_sccs).

        issue_rows — list of (src_label, src_type, issue_text, src_pts, dst_pts)
        n_sccs     — total number of strongly-connected components found
        """
        names = set(all_points.keys())
        # Keep only edges between known points
        flt_out = {v: outgoing[v] & names for v in names if v in outgoing}

        # ── Tarjan's iterative SCC ──────────────────────────────────────────
        idx_map  = {}
        lowlink  = {}
        on_stk   = set()
        stk      = []
        ctr      = [0]
        sccs     = []

        def _visit(root):
            work = [(root, iter(flt_out.get(root, set())))]
            idx_map[root] = lowlink[root] = ctr[0]; ctr[0] += 1
            stk.append(root); on_stk.add(root)
            while work:
                v, nbrs = work[-1]
                try:
                    w = next(nbrs)
                    if w not in idx_map:
                        idx_map[w] = lowlink[w] = ctr[0]; ctr[0] += 1
                        stk.append(w); on_stk.add(w)
                        work.append((w, iter(flt_out.get(w, set()))))
                    elif w in on_stk:
                        lowlink[v] = min(lowlink[v], idx_map[w])
                except StopIteration:
                    work.pop()
                    if work:
                        lowlink[work[-1][0]] = min(lowlink[work[-1][0]], lowlink[v])
                    if lowlink[v] == idx_map[v]:
                        scc = []
                        while True:
                            w = stk.pop(); on_stk.discard(w)
                            scc.append(w)
                            if w == v:
                                break
                        sccs.append(frozenset(scc))

        for n in sorted(names):
            if n not in idx_map:
                _visit(n)

        if len(sccs) <= 1:
            return [], len(sccs)

        # ── Build condensation graph ────────────────────────────────────────
        node_scc = {}
        for i, scc in enumerate(sccs):
            for n in scc:
                node_scc[n] = i

        n_s = len(sccs)
        cond = [set() for _ in range(n_s)]
        for v in names:
            sv = node_scc[v]
            for w in flt_out.get(v, set()):
                sw = node_scc.get(w)
                if sw is not None and sv != sw:
                    cond[sv].add(sw)

        # ── BFS reachability on condensation ───────────────────────────────
        reachable = []
        for i in range(n_s):
            vis = set()
            q = deque([i])
            while q:
                cur = q.popleft()
                for nxt in cond[cur]:
                    if nxt not in vis:
                        vis.add(nxt)
                        q.append(nxt)
            reachable.append(vis)

        # ── Emit one row per unreachable (src_scc → dst_scc) direction ─────
        rows = []
        for i in range(n_s):
            for j in range(n_s):
                if i == j or j in reachable[i]:
                    continue
                src_pts = sorted(sccs[i])
                dst_pts = sorted(sccs[j])
                src_rep   = src_pts[0]
                src_type  = all_points.get(src_rep, "?")
                src_label = (src_rep if len(src_pts) == 1
                             else f"{src_rep} (grp {len(src_pts)})")
                preview = ", ".join(dst_pts[:4])
                if len(dst_pts) > 4:
                    preview += f"  +{len(dst_pts) - 4} more"
                issue = f"Cannot reach {len(dst_pts)} pt(s): {preview}"
                rows.append((src_label, src_type, issue, src_pts, dst_pts))

        return rows, n_s

    def _run_map_validation(self):
        path = self._tab4_file_entry.get().strip()
        if not path or not os.path.exists(path):
            messagebox.showerror("Error", "Select a valid .json or .smap file first")
            return

        # Clear previous results
        for item in self._tab4_tree.get_children():
            self._tab4_tree.delete(item)
        self._tab4_log.clear()
        for w in self._tab4_summary.winfo_children():
            w.destroy()

        self._tab4_log.log("Loading file…", "info")
        self.update_status("Map Validator: loading…")
        try:
            data = self._safe_load_json(path)
            if data is None:
                return
        except Exception as e:
            self._tab4_log.log(str(e), "err")
            return

        points = data.get("advancedPointList", [])
        curves = data.get("advancedCurveList", [])

        # Build incoming / outgoing sets per node name
        outgoing = defaultdict(set)   # node → set of reachable neighbours
        incoming = defaultdict(set)   # node → set of nodes that reach it

        for curve in curves:
            src = (curve.get("startPos") or {}).get("instanceName")
            dst = (curve.get("endPos")   or {}).get("instanceName")
            if src and dst and src != dst:
                outgoing[src].add(dst)
                incoming[dst].add(src)

        # Collect all named points
        all_points = {}
        for pt in points:
            name = pt.get("instanceName")
            cls  = pt.get("className", "Unknown")
            if name:
                all_points[name] = "AP" if cls == "ActionPoint" else "LM"

        issues = []   # (name, pt_type, issue_text)
        for name, pt_type in sorted(all_points.items()):
            has_in  = bool(incoming.get(name))
            has_out = bool(outgoing.get(name))
            if not has_in and not has_out:
                issues.append((name, pt_type, "No entry path  &  No exit path"))
            elif not has_in:
                issues.append((name, pt_type, "No entry path"))
            elif not has_out:
                issues.append((name, pt_type, "No exit path"))

        # ── Check 1: entry / exit paths ────────────────────────────────────
        for name, pt_type, issue in issues:
            tag = "warn" if "exit" in issue.lower() else "err"
            self._tab4_tree.insert("", "end", values=(name, pt_type, issue), tags=(tag,))

        if issues:
            self._tab4_log.log(f"Found {len(issues)} entry/exit issue(s).", "warn")
            no_entry = sum(1 for _, _, i in issues if "No entry" in i)
            no_exit  = sum(1 for _, _, i in issues if "No exit"  in i)
            both     = sum(1 for _, _, i in issues if "&"        in i)
            self._tab4_log.log(
                f"  No entry only: {no_entry - both}  |  "
                f"No exit only: {no_exit - both}  |  "
                f"Both missing: {both}", "warn")
        else:
            self._tab4_log.log(
                f"✅ All {len(all_points)} points have valid entry and exit paths.", "ok")

        # ── Check 2: global reachability (SCC + condensation BFS) ──────────
        self._tab4_log.log("Checking global reachability…", "info")
        conn_rows, n_sccs = self._check_connectivity(all_points, outgoing)
        n_conn = len(conn_rows)

        if conn_rows:
            # separator row
            self._tab4_tree.insert("", "end",
                values=("── Connectivity ──", "",
                        f"{n_conn} unreachable direction(s)  |  {n_sccs} group(s)"),
                tags=("conn",))
            for src_label, src_type, issue, src_pts, dst_pts in conn_rows:
                self._tab4_tree.insert("", "end",
                    values=(src_label, src_type, issue), tags=("conn",))
            self._tab4_log.log(
                f"⚠ Graph has {n_sccs} disconnected group(s), "
                f"{n_conn} unreachable direction(s).", "warn")
            for src_label, src_type, issue, src_pts, dst_pts in conn_rows:
                src_str = ", ".join(src_pts)
                dst_str = ", ".join(dst_pts)
                self._tab4_log.log(
                    f"  [{src_str}]  →✗→  [{dst_str}]", "warn")
        else:
            self._tab4_log.log("✅ All points are mutually reachable.", "ok")

        # ── Summary bar ────────────────────────────────────────────────────
        total   = len(all_points)
        n_issue = len(issues)
        n_ok    = total - n_issue
        for txt, color in [
            (f"✅  {n_ok} OK",                              SUCCESS),
            (f"⚠️  {n_issue} path issue(s)",               WARNING if n_issue else SUCCESS),
            (f"🔗  {n_conn} connectivity gap(s)",          WARNING if n_conn  else SUCCESS),
            (f"📍  {total} total points",                   TEXT_MID),
        ]:
            ctk.CTkLabel(self._tab4_summary, text=txt,
                         font=("Segoe UI", 11, "bold"),
                         text_color=color).pack(side="left", padx=(0, 20))

        self.update_status(
            f"Map Validator: {n_issue} path issue(s), {n_conn} connectivity gap(s)")

    # ─────────────────── ALL BUSINESS LOGIC (original, untouched) ───────────────────

    def on_mode_change(self, *args):
        mode = self.mode_var.get()
        if mode == "automatic":
            self.loading_combo.configure(state="readonly")
        else:
            self.loading_combo.configure(state="disabled")

    # ── Info popups ──
    def show_json_format_info(self):
        messagebox.showinfo(
            "JSON Format Required",
            "Your JSON must contain these sections:\n\n"
            "1) advancedPointList: ActionPoint / LocationMark entries\n"
            "   Example ActionPoint:\n"
            "   {\"className\":\"ActionPoint\",\"instanceName\":\"AP123\",\"pos\":{\"x\":10.0,\"y\":20.0}}\n\n"
            "2) advancedCurveList: path entries (DegenerateBezier)\n"
            "   Example curve:\n"
            "   {\"className\":\"DegenerateBezier\",\"startPos\":{\"instanceName\":\"LM1\",\"pos\":{\"x\":...}},"
            "\"endPos\":{\"instanceName\":\"AP2\",\"pos\":{\"x\":...}}}\n\n"
            "Missing curves or missing x/y will break path calculations."
        )

    def show_marker_json_info(self):
        messagebox.showinfo(
            "Markers JSON Required",
            "Your JSON for marker extraction should contain a 'markers' array with entries like:\n"
            "{\"type\":\"WorkMarker\",\"code\":\"M1\",\"markInfos\":[{\"x\":10.0,\"y\":20.0}, ...]}\n\n"
            "Ensure markInfos contain x and y coordinates."
        )

    def show_excel_format_info(self):
        if self.ndeep_enabled_var and self.ndeep_enabled_var.get():
            messagebox.showinfo(
                "Excel Format (N-Deep mode)",
                "In N-Deep mode the Excel must contain EXACT columns:\n"
                "Zone name | zone entry | zone exit | min x | max x | min y | max y\n\n"
                "Subzone columns are NOT required — subzones are generated automatically."
            )
        else:
            messagebox.showinfo(
                "Excel Format Required",
                "Excel must contain EXACT columns:\n"
                "Zone name | zone entry | zone exit | subzone name | subzone entery | subzone exit | min x | max x | min y | max y"
            )

    # ── Safe JSON load ──
    def _safe_load_json(self, path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if not data:
                messagebox.showerror("Error", "JSON file is empty")
                return None
            return data
        except json.JSONDecodeError as e:
            messagebox.showerror("JSON Parse Error", f"Invalid JSON: {e}")
            return None
        except Exception as e:
            messagebox.showerror("File Error", f"Cannot read JSON: {e}")
            return None

    # ── Validations ──
    def validate_json_for_pick_sequence(self, data):
        if not isinstance(data, dict):
            return False, "JSON root must be an object"
        if "advancedPointList" not in data:
            return False, "Missing 'advancedPointList'"
        if "advancedCurveList" not in data:
            return False, "Missing 'advancedCurveList'"
        ap_found = False
        for p in data.get("advancedPointList", []):
            if p.get("className") == "ActionPoint":
                pos = p.get("pos") or {}
                if "x" in pos and "y" in pos:
                    ap_found = True
                    break
        if not ap_found:
            return False, "No ActionPoint with valid pos found in 'advancedPointList'"
        for c in data.get("advancedCurveList", []):
            s = c.get("startPos", {})
            e = c.get("endPos", {})
            if not s.get("instanceName") or not e.get("instanceName"):
                return False, "A curve entry is missing start/end instanceName"
            if "pos" not in s or "pos" not in e:
                return False, "A curve entry is missing position ('pos') for start or end"
            if "x" not in s.get("pos", {}) or "y" not in s.get("pos", {}):
                return False, "Start position missing x/y in a curve"
            if "x" not in e.get("pos", {}) or "y" not in e.get("pos", {}):
                return False, "End position missing x/y in a curve"
        return True, "OK"

    def validate_json_for_markers(self, data):
        if not isinstance(data, dict):
            return False, "Markers JSON must be an object"
        if "markers" not in data:
            return False, "Missing 'markers' array in JSON"
        markers = data.get("markers", [])
        if not isinstance(markers, list) or not markers:
            return False, "'markers' must be a non-empty list"
        for m in markers:
            if not isinstance(m, dict):
                continue
            infos = m.get("markInfos", [])
            if isinstance(infos, list):
                for info in infos:
                    if isinstance(info, dict) and "x" in info and "y" in info:
                        return True, "OK"
        return False, "No valid markInfos with x/y found in markers"

    def validate_excel_for_zone(self, df):
        required = ['Zone name', 'zone entry', 'zone exit', 'subzone name',
                    'subzone entery', 'subzone exit', 'min x', 'max x', 'min y', 'max y']
        missing = [c for c in required if c not in df.columns]
        if missing:
            return False, f"Missing columns: {missing}"
        try:
            for col in ['min x', 'max x', 'min y', 'max y']:
                pd.to_numeric(df[col], errors='raise')
        except Exception as e:
            return False, f"Coordinate columns must be numeric: {e}"
        return True, "OK"

    # ── Tab1 logic ──
    def run_marker_extraction(self):
        input_file = self._tab1_file_picker.get()
        output_file = self.tab1_output_entry.get()
        if not input_file or not os.path.exists(input_file):
            messagebox.showerror("Error", "Please select a valid JSON file")
            return
        try:
            self.update_status("Validating JSON for markers...")
            self.log_tab1("🔍 Validating JSON file...")
            data = self._safe_load_json(input_file)
            if data is None:
                self.update_status("JSON parse error")
                return
            ok, msg = self.validate_json_for_markers(data)
            if not ok:
                messagebox.showerror("Invalid JSON (Markers)", msg)
                self.log_tab1("❌ " + msg)
                self.update_status("Validation failed")
                return
            self.log_tab1("✅ Marker JSON validated")
            self.update_status("Extracting markers...")
            work_sheets = []
            validation_data = {}
            if self.work_marker_var.get():
                self.log_tab1("🔄 Processing WorkMarkers...")
                work_markers = self.extract_markers(data, "WorkMarker")
                work_result, missing = self.process_markers(work_markers)
                work_sheets.append(("WorkMarkers_Grouped", work_result))
                if missing:
                    validation_data["Missing_WorkMarkers"] = list(missing)
                self.log_tab1(f"Found {len(work_markers)} WorkMarkers")
            if self.resource_marker_var.get():
                self.log_tab1("🔄 Processing ResourceMarkers...")
                res_markers = self.extract_markers(data, "ResourceMarker")
                res_result, missing = self.process_markers(res_markers)
                work_sheets.append(("ResourceMarkers_Grouped", res_result))
                if missing:
                    validation_data["Missing_ResourceMarkers"] = list(missing)
                self.log_tab1(f"Found {len(res_markers)} ResourceMarkers")
            with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
                for sheet_name, df in work_sheets:
                    if isinstance(df, pd.DataFrame) and df.empty:
                        pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)
                    else:
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                for sheet, missing_codes in validation_data.items():
                    pd.DataFrame({"Missing Codes": missing_codes}).to_excel(
                        writer, sheet_name=sheet, index=False)
            messagebox.showinfo("Success", f"Saved to {output_file}")
            self.update_status("Marker extraction complete!")
            self.log_tab1("✅ Extraction saved")
        except Exception as e:
            self.log_tab1(f"💥 Error: {e}")
            messagebox.showerror("Error", str(e))
            self.update_status("Error during extraction")

    def extract_markers(self, data, marker_type):
        return [
            {"code": m.get("code"), "x": info.get("x"), "y": info.get("y")}
            for m in data.get("markers", []) if m.get("type") == marker_type
            for info in m.get("markInfos", [])
            if info.get("x") is not None and info.get("y") is not None
        ]

    def process_markers(self, markers):
        df = pd.DataFrame(markers)
        if df.empty:
            return pd.DataFrame(), set()
        grouped = {}
        for x, group in df.groupby("x"):
            sorted_group = group.sort_values("y")
            grouped[x] = sorted_group["code"].tolist()
        max_len = max(len(codes) for codes in grouped.values())
        aligned = {x: codes + [None] * (max_len - len(codes)) for x, codes in grouped.items()}
        result_df = pd.DataFrame(aligned)
        all_codes = set(df["code"])
        used_codes = set(code for lst in grouped.values() for code in lst if code)
        missing_codes = all_codes - used_codes
        return result_df, missing_codes

    # ── Tab2 logic ──
    def load_and_analyze(self):
        input_file = self._tab2_file_entry.get()
        if not input_file or not os.path.exists(input_file):
            messagebox.showerror("Error", "Select a valid JSON file")
            return
        try:
            self.update_status("Validating JSON for pick sequence...")
            self.log_tab2("🔍 Validating JSON file...")
            data = self._safe_load_json(input_file)
            if data is None:
                return
            ok, msg = self.validate_json_for_pick_sequence(data)
            if not ok:
                messagebox.showerror("Invalid JSON (Pick Sequence)", msg)
                self.log_tab2("❌ " + msg)
                self.update_status("Validation failed")
                return
            self.log_tab2("✅ JSON structure valid for pick sequence")
            self.update_status("Building graph...")
            self.json_data = data
            self.graph, self.node_positions = self.build_graph_from_json(self.json_data)
            self.log_tab2(f"🔗 Graph built: {len(self.graph)} nodes")
            aps = self.get_action_points(self.json_data)
            if not aps:
                messagebox.showerror("Error", "No ActionPoints found")
                return
            excluded_aps = list(
                self._map_excluded | self._map_loading_aps | self._map_unloading_aps)
            self.excluded_aps = excluded_aps
            if excluded_aps:
                self.log_tab2(f"🚫 Excluding {len(excluded_aps)} APs "
                              f"({len(self._map_loading_aps)} load, "
                              f"{len(self._map_unloading_aps)} unload, "
                              f"{len(self._map_excluded)} removed) from pick sequence")
            aps = [ap for ap in aps if ap['name'] not in excluded_aps]
            orientation = self.orientation_var.get()
            self.log_tab2(f"📊 Grouping aisles ({orientation})...")
            self.aisles, axis = self.auto_group_aisles(aps, orientation)
            self.aisles = [a for a in self.aisles if a]
            if self.cross_aisle_var.get():
                self.log_tab2("🔍 Detecting cross aisles...")
                self.aisles = self.detect_all_cross_aisles(self.aisles, axis)
            if self._map_cross_aisle_lms:
                self.log_tab2(
                    f"✚ Splitting by {len(self._map_cross_aisle_lms)} manual cross-aisle marker(s)...")
                self.aisles = self._split_aisles_by_markers(self.aisles, axis)
            filtered_aisles = []
            for aisle in self.aisles:
                filtered = [ap for ap in aisle if ap['name'] in self.node_positions and ap['name'] not in excluded_aps]
                if filtered:
                    filtered_aisles.append(filtered)
            self.aisles = filtered_aisles
            if not self.aisles:
                messagebox.showwarning("Warning", "No aisles remain after filtering/exclusions")
                self.log_tab2("⚠️ No aisles remain after filtering/exclusions")
            ap_names = [ap['name'] for ap in aps if ap['name'] in self.node_positions and ap['name'] not in excluded_aps]
            self.loading_combo.configure(values=ap_names)
            if ap_names:
                self.loading_combo.set(ap_names[0])
            if ap_names:
                lp = ap_names[0]
                dist, _ = self.dijkstra(lp) if lp in self.graph else ({}, {})
                unreachable = [ap for ap in ap_names if dist.get(ap, float('inf')) == float('inf')]
                if unreachable:
                    self.log_tab2(f"⚠️ Warning: {len(unreachable)} APs are unreachable from default loading point ({lp}).")
                    messagebox.showwarning("Warning",
                                           f"{len(unreachable)} AP(s) appear unreachable from loading point {lp}. Check curves or choose another loading point.")
            self.display_aisles()
            self.tab2_ap_count_label.configure(text=str(len(ap_names)))
            self.log_tab2(f"✅ Analysis complete — {len(self.aisles)} aisles, {len(ap_names)} APs")
            self.update_status("Aisle analysis complete")
        except Exception as e:
            self.log_tab2(str(e))
            messagebox.showerror("Error", str(e))
            self.update_status("Error during analysis")

    def build_graph_from_json(self, json_data):
        graph = {}
        node_positions = {}
        for p in json_data.get("advancedPointList", []):
            inst = p.get("instanceName")
            pos = p.get("pos", {})
            if inst and "x" in pos and "y" in pos:
                node_positions[inst] = (float(pos["x"]), float(pos["y"]))
                graph.setdefault(inst, {})
        for curve in json_data.get("advancedCurveList", []):
            try:
                start = curve.get("startPos", {}).get("instanceName")
                end = curve.get("endPos", {}).get("instanceName")
                sx = curve.get("startPos", {}).get("pos", {}).get("x")
                sy = curve.get("startPos", {}).get("pos", {}).get("y")
                ex = curve.get("endPos", {}).get("pos", {}).get("x")
                ey = curve.get("endPos", {}).get("pos", {}).get("y")
                if start and (start not in node_positions) and sx is not None and sy is not None:
                    node_positions[start] = (float(sx), float(sy))
                    graph.setdefault(start, {})
                if end and (end not in node_positions) and ex is not None and ey is not None:
                    node_positions[end] = (float(ex), float(ey))
                    graph.setdefault(end, {})
            except Exception:
                continue
        for curve in json_data.get("advancedCurveList", []):
            start = curve.get("startPos", {}).get("instanceName")
            end = curve.get("endPos", {}).get("instanceName")
            if not start or not end:
                continue
            s_pos = curve.get("startPos", {}).get("pos", {})
            e_pos = curve.get("endPos", {}).get("pos", {})
            if ("x" not in s_pos) or ("y" not in s_pos) or ("x" not in e_pos) or ("y" not in e_pos):
                continue
            sx, sy = float(s_pos["x"]), float(s_pos["y"])
            ex, ey = float(e_pos["x"]), float(e_pos["y"])
            weight = math.hypot(ex - sx, ey - sy)
            if weight <= 0:
                continue
            graph.setdefault(start, {})[end] = min(weight, graph.get(start, {}).get(end, weight))
        return graph, node_positions

    def dijkstra(self, source, graph=None):
        if graph is None:
            graph = self.graph
        if source not in graph:
            return {}, {}
        dist = {n: float('inf') for n in graph}
        prev = {n: None for n in graph}
        dist[source] = 0.0
        heap = [(0.0, source)]
        while heap:
            d, u = heapq.heappop(heap)
            if d > dist[u]:
                continue
            for v, w in graph.get(u, {}).items():
                nd = d + w
                if nd < dist.get(v, float('inf')):
                    dist[v] = nd
                    prev[v] = u
                    heapq.heappush(heap, (nd, v))
        return dist, prev

    def calculate_distance(self, a, b):
        if isinstance(a, dict) and 'name' in a:
            a_name = a['name']
        else:
            a_name = a
        if isinstance(b, dict) and 'name' in b:
            b_name = b['name']
        else:
            b_name = b
        if a_name == b_name:
            return 0.0
        if a_name not in self.graph or b_name not in self.graph:
            pos_a = self.node_positions.get(a_name)
            pos_b = self.node_positions.get(b_name)
            if pos_a and pos_b:
                return math.hypot(pos_a[0] - pos_b[0], pos_a[1] - pos_b[1])
            return float('inf')
        dist, _ = self.dijkstra(a_name)
        return dist.get(b_name, float('inf'))

    def optimize_aisle_order(self, aisles, loading_point, orientation):
        if not aisles:
            return []
        lp_name = loading_point['name'] if isinstance(loading_point, dict) else loading_point
        dist_from_lp, _ = self.dijkstra(lp_name) if lp_name in self.graph else ({}, {})
        aisle_scores = []
        for aisle in aisles:
            min_d = float('inf')
            for ap in aisle:
                ap_name = ap['name']
                d = dist_from_lp.get(ap_name, float('inf')) if dist_from_lp else self.calculate_distance(lp_name, ap_name)
                if d < min_d:
                    min_d = d
            aisle_scores.append((min_d, aisle))
        aisle_scores.sort(key=lambda x: (math.isinf(x[0]), x[0]))
        return [a for _, a in aisle_scores]

    def optimize_aisle_direction(self, aisle, loading_point, prev_aisle_end=None):
        first_ap = aisle[0]
        last_ap = aisle[-1]
        if prev_aisle_end is None:
            d_first = self.calculate_distance(loading_point, first_ap)
            d_last = self.calculate_distance(loading_point, last_ap)
            if d_first <= d_last:
                return aisle, last_ap
            else:
                return list(reversed(aisle)), first_ap
        else:
            d_first = self.calculate_distance(prev_aisle_end, first_ap)
            d_last = self.calculate_distance(prev_aisle_end, last_ap)
            if d_first <= d_last:
                return aisle, last_ap
            else:
                return list(reversed(aisle)), first_ap

    def generate_automatic_sequence(self, aisles, loading_point, orientation):
        ordered_aisles = self.optimize_aisle_order(aisles, loading_point, orientation)
        final_sequence = []
        current_position = loading_point
        for aisle in ordered_aisles:
            aisle_nodes = [ap for ap in aisle if ap['name'] in self.node_positions]
            if not aisle_nodes:
                self.log_tab2(f"⚠️ Skipping aisle (no graph nodes): {[ap['name'] for ap in aisle]}")
                continue
            optimized_aisle, end_point = self.optimize_aisle_direction(aisle_nodes, loading_point, current_position)
            start_ap = optimized_aisle[0]
            dist_to_start = self.calculate_distance(current_position, start_ap)
            if math.isinf(dist_to_start):
                self.log_tab2(f"⚠️ Skipping aisle (unreachable from current pos): {start_ap['name']}")
                continue
            final_sequence.extend(optimized_aisle)
            current_position = end_point
        return final_sequence

    def display_aisles(self):
        for item in self.aisle_tree.get_children():
            self.aisle_tree.delete(item)
        for i, aisle in enumerate(self.aisles, 1):
            first = aisle[0]["name"]
            last  = aisle[-1]["name"]
            count = len({ap["name"] for ap in aisle})
            self.aisle_tree.insert("", "end",
                values=(f"Aisle {i}", first, last, count, "First→Last"))
        self.aisle_tree.bind("<Double-1>", self.on_direction_click)

    def on_direction_click(self, event):
        item = self.aisle_tree.identify_row(event.y)
        col = self.aisle_tree.identify_column(event.x)
        if col == "#5":
            current = self.aisle_tree.item(item, "values")[4]
            new = "Last→First" if current == "First→Last" else "First→Last"
            self.aisle_tree.set(item, "Direction", new)

    def set_all_directions(self, direction):
        text = "First→Last" if direction == "f2l" else "Last→First"
        for item in self.aisle_tree.get_children():
            self.aisle_tree.set(item, "Direction", text)

    def toggle_selected_directions(self):
        for item in self.aisle_tree.selection():
            cur = self.aisle_tree.item(item, "values")[4]
            new = "Last→First" if cur == "First→Last" else "First→Last"
            self.aisle_tree.set(item, "Direction", new)

    def generate_sequence(self):
        if not self.aisles:
            messagebox.showwarning("Warning", "Load & analyze first")
            return
        filename = self.tab2_output_entry.get()
        dup = 0
        if self.dup_enabled_var.get():
            try:
                dup = int(self.dup_count_entry.get())
            except ValueError:
                messagebox.showerror("Error", "Duplication count must be a whole number")
                return
            if dup < 2:
                messagebox.showerror("Error", "Duplication count must be at least 2")
                return
        template_path = self._tab2_template_entry.get()
        try:
            if self.mode_var.get() == "automatic":
                self.log_tab2("🤖 Generating automatic sequence...")
                loading_point_name = self.loading_point_var.get()
                if not loading_point_name:
                    messagebox.showerror("Error", "Please select a Loading Point AP")
                    return
                loading_point = None
                for ap in self.get_action_points(self.json_data):
                    if ap['name'] == loading_point_name:
                        loading_point = ap
                        break
                if not loading_point:
                    messagebox.showerror("Error", "Loading point not found in AP list")
                    return
                if loading_point_name not in self.node_positions and loading_point_name not in self.graph:
                    messagebox.showerror("Error", "Selected loading point has no coordinates or graph node")
                    return
                sequence = self.generate_automatic_sequence(self.aisles, loading_point, self.orientation_var.get())
                if not sequence:
                    messagebox.showwarning("Warning", "Generated sequence is empty (check graph/connectivity/exclusions)")
                self.save_sequence_to_excel(sequence, filename, automatic_mode=True, duplication=dup, template_path=template_path)
                self.log_tab2(f"Done! {len(sequence)} APs written" + (f" + duplicated (x{dup})" if dup else ""))
                self.update_status("Sequence generated")
            else:
                self.log_tab2("👤 Generating manual sequence...")
                directions = []
                for item in self.aisle_tree.get_children():
                    direction = self.aisle_tree.item(item, "values")[4]
                    directions.append("f2l" if direction == "First→Last" else "l2f")
                final = []
                excluded = getattr(self, "excluded_aps", []) or []
                for aisle, d in zip(self.aisles, directions):
                    clean_aisle = [ap for ap in aisle if ap['name'] not in excluded]
                    if not clean_aisle:
                        continue
                    final.append(clean_aisle if d == "f2l" else list(reversed(clean_aisle)))
                self.save_sequence_to_excel(final, filename, automatic_mode=False, duplication=dup, template_path=template_path)
                self.update_status("Manual sequence saved")
                messagebox.showinfo("Success", f"Saved to {filename}")
        except Exception as e:
            self.log_tab2(str(e))
            messagebox.showerror("Error", str(e))
            self.update_status("Error generating sequence")

    def get_action_points(self, json_data):
        aps = []
        for ap in json_data.get("advancedPointList", []):
            if ap.get("className") == "ActionPoint":
                pos = ap.get("pos", {})
                if "x" in pos and "y" in pos:
                    aps.append({"name": ap.get("instanceName"), "x": float(pos["x"]), "y": float(pos["y"])})
        return aps

    def auto_group_aisles(self, aps, orientation="horizontal", gap=0.2):
        if not aps:
            return [], "x"
        if orientation == "horizontal":
            aps.sort(key=lambda ap: (round(ap["y"], 6), ap["x"]))
            axis = "y"
            sort_axis = "x"
        else:
            aps.sort(key=lambda ap: (round(ap["x"], 6), ap["y"]))
            axis = "x"
            sort_axis = "y"
        aisles = []
        current = [aps[0]]
        last = aps[0][axis]
        for ap in aps[1:]:
            if abs(ap[axis] - last) > gap:
                aisles.append(current)
                current = []
            current.append(ap)
            last = ap[axis]
        aisles.append(current)
        for aisle in aisles:
            aisle.sort(key=lambda ap: ap[sort_axis])
        return aisles, sort_axis

    def _split_aisles_by_markers(self, aisles, sort_axis):
        """
        Split aisles at positions defined by manually-tagged cross-aisle markers.
        For each tagged point, its sort_axis coordinate acts as a cut line —
        any aisle whose APs straddle that coordinate gets split there.
        """
        cross_positions = sorted(
            pt[sort_axis] for pt in self._map_points
            if pt["name"] in self._map_cross_aisle_lms
        )
        if not cross_positions:
            return aisles

        output = []
        for aisle in aisles:
            aisle_sorted = sorted(aisle, key=lambda ap: ap[sort_axis])
            coords = [ap[sort_axis] for ap in aisle_sorted]

            split_indices = set()
            for cp in cross_positions:
                for i in range(len(coords) - 1):
                    lo = min(coords[i], coords[i + 1])
                    hi = max(coords[i], coords[i + 1])
                    if lo < cp < hi:
                        split_indices.add(i + 1)
                        break

            if not split_indices:
                output.append(aisle_sorted)
                continue

            start = 0
            for idx in sorted(split_indices):
                section = aisle_sorted[start:idx]
                if section:
                    output.append(section)
                start = idx
            last = aisle_sorted[start:]
            if last:
                output.append(last)

        return output

    def detect_all_cross_aisles(self, aisles, sort_axis, sensitivity=1.5):
        output = []
        for aisle in aisles:
            if len(aisle) >= 3:
                subs, _ = self.detect_cross_aisles(aisle, sort_axis, sensitivity)
                output.extend(subs)
            else:
                output.append(aisle)
        return output

    def detect_cross_aisles(self, aisle, sort_axis, sensitivity=1.5, min_length=3):
        if len(aisle) < min_length:
            return [aisle], []
        distances = [abs(aisle[i + 1][sort_axis] - aisle[i][sort_axis]) for i in range(len(aisle) - 1)]
        if not distances:
            return [aisle], []
        arr = np.array(distances)
        q25, q50, q75 = np.percentile(arr, [25, 50, 75])
        iqr = q75 - q25
        if iqr > 0:
            threshold = q75 + sensitivity * iqr
        else:
            threshold = q50 * 2.5
        threshold = max(threshold, q50 * 1.8)
        splits = sorted(i for i, d in enumerate(distances) if d > threshold)
        if not splits:
            return [aisle], []
        # Build section boundaries and check each resulting section has >= 2 APs.
        # Using per-section sizes (not cumulative left/right) so that multiple
        # cross-aisles don't incorrectly fail the size constraint.
        boundaries = [0] + [s + 1 for s in splits] + [len(aisle)]
        section_sizes = [boundaries[k + 1] - boundaries[k]
                         for k in range(len(boundaries) - 1)]
        validated = [
            s for s, (sz_l, sz_r) in zip(splits, zip(section_sizes, section_sizes[1:]))
            if sz_l >= 2 and sz_r >= 2
        ]
        if not validated:
            return [aisle], []
        idxs = [i + 1 for i in validated]
        subaisles = []
        start = 0
        for idx in idxs:
            subaisles.append(aisle[start:idx])
            start = idx
        subaisles.append(aisle[start:])
        return subaisles, []

    def save_sequence_to_excel(self, sequence, filename, automatic_mode=False,
                               duplication=0, template_path=""):
        # Flatten to ordered list of AP names
        if automatic_mode:
            flat = [ap["name"] for ap in sequence]
        else:
            flat = [ap["name"] for aisle in sequence for ap in aisle]

        if template_path and os.path.exists(template_path):
            wb = load_workbook(template_path)

            # ── Sheet 1: location-marker-mapping (duplication) ──
            ws_dup = wb["location-marker-mapping"]
            # Clear all data rows, keep header
            for r in range(2, ws_dup.max_row + 1):
                for c in range(1, ws_dup.max_column + 1):
                    ws_dup.cell(r, c).value = None
            dup_count = duplication if duplication >= 2 else 1
            row = 2
            for ap_name in flat:
                for _ in range(dup_count):
                    ws_dup.cell(row, 1).value = None   # Location
                    ws_dup.cell(row, 2).value = ap_name  # Marker Code
                    ws_dup.cell(row, 3).value = False    # Is Smart
                    row += 1

            # ── Sheet 2: marker-sequence (pick sequence) ──
            ws_seq = wb["marker-sequence"]
            for r in range(2, ws_seq.max_row + 1):
                for c in range(1, ws_seq.max_column + 1):
                    ws_seq.cell(r, c).value = None
            for i, ap_name in enumerate(flat, 2):
                ws_seq.cell(i, 1).value = None     # Seq
                ws_seq.cell(i, 2).value = ap_name  # Marker Code

            # ── Sheet 5: station-config (template path) ──
            ws_sc = wb["station-config"]
            # Clear old data rows
            for r in range(2, ws_sc.max_row + 1):
                for col in range(1, ws_sc.max_column + 1):
                    ws_sc.cell(r, col).value = None
            self._write_station_config_rows(ws_sc)

            wb.save(filename)
        else:
            # No template: build all 5 sheets with standard headers
            wb = Workbook()

            # Sheet 1: location-marker-mapping (duplication)
            ws_dup = wb.active
            ws_dup.title = "location-marker-mapping"
            ws_dup.cell(1, 1, "Location")
            ws_dup.cell(1, 2, "Marker Code")
            ws_dup.cell(1, 3, "Is Smart")
            dup_count = duplication if duplication >= 2 else 1
            row = 2
            for ap_name in flat:
                for _ in range(dup_count):
                    ws_dup.cell(row, 1).value = None
                    ws_dup.cell(row, 2).value = ap_name
                    ws_dup.cell(row, 3).value = False
                    row += 1

            # Sheet 2: marker-sequence (pick sequence)
            ws_seq = wb.create_sheet("marker-sequence")
            ws_seq.cell(1, 1, "Seq")
            ws_seq.cell(1, 2, "Marker Code")
            for i, ap_name in enumerate(flat, 2):
                ws_seq.cell(i, 2, ap_name)

            # Sheet 3: aisle-sequence
            ws_aisle = wb.create_sheet("aisle-sequence")
            ws_aisle.cell(1, 1, "Seq")
            ws_aisle.cell(1, 2, "Aisle")

            # Sheet 4: marker-config
            ws_mc = wb.create_sheet("marker-config")
            for col, h in enumerate(
                    ["MarkerId", "Name", "Type", "Queue Capacity", "Is Active"], 1):
                ws_mc.cell(1, col, h)

            # Sheet 5: station-config
            ws_sc = wb.create_sheet("station-config")
            for col, h in enumerate(
                    ["stationGroupUid", "prePoint", "stationUid", "labels",
                     "stationType", "stationMarker", "waitMarker", "priority",
                     "active", "capacity", "height", "Ownership"], 1):
                ws_sc.cell(1, col, h)
            self._write_station_config_rows(ws_sc)

            wb.save(filename)

    def _write_station_config_rows(self, ws):
        """Write one row per loading AP then one per unloading AP to ws."""
        row = 2
        for group, label, ap_set in [
            ("LOADING",   "loading",   sorted(self._map_loading_aps)),
            ("UNLOADING", "unloading", sorted(self._map_unloading_aps)),
        ]:
            for i, ap_name in enumerate(ap_set, 1):
                wait_marker = self._find_closest_lm(ap_name)
                ws.cell(row, 1,  group)
                ws.cell(row, 2,  None)                   # prePoint — blank
                ws.cell(row, 3,  f"{group}_{i}")         # stationUid
                ws.cell(row, 4,  label)                  # labels
                ws.cell(row, 5,  "stp")                  # stationType
                ws.cell(row, 6,  ap_name)                # stationMarker
                ws.cell(row, 7,  wait_marker)            # waitMarker
                ws.cell(row, 8,  1)                      # priority
                ws.cell(row, 9,  True)                   # active
                ws.cell(row, 10, 1)                      # capacity
                ws.cell(row, 11, 8)                      # height
                ws.cell(row, 12, "SELF")                 # Ownership
                row += 1

    def _find_closest_lm(self, ap_name: str):
        """Return the name of the LM closest to ap_name via graph path,
        falling back to Euclidean distance if no path exists."""
        if not self.json_data:
            return None
        lm_positions = {}
        for pt in self.json_data.get("advancedPointList", []):
            if pt.get("className") != "ActionPoint":
                pos = pt.get("pos", {})
                inst = pt.get("instanceName")
                if inst and "x" in pos and "y" in pos:
                    lm_positions[inst] = (float(pos["x"]), float(pos["y"]))
        if not lm_positions:
            return None
        # Graph-based shortest path
        if ap_name in self.graph:
            dist, _ = self.dijkstra(ap_name)
            reachable = {n: d for n, d in dist.items()
                         if n in lm_positions and d < float("inf")}
            if reachable:
                return min(reachable, key=reachable.get)
        # Euclidean fallback
        if ap_name not in self.node_positions:
            return None
        ax, ay = self.node_positions[ap_name]
        return min(lm_positions,
                   key=lambda n: math.hypot(lm_positions[n][0] - ax,
                                            lm_positions[n][1] - ay))

    # ─────────────────── Map helpers (used by Pick Sequence tab) ───────────────────

    def _map_fit_all(self):
        if not self._map_points:
            return
        xs = [p["x"] for p in self._map_points]
        ys = [p["y"] for p in self._map_points]
        self._map_world_cx = (min(xs) + max(xs)) / 2
        self._map_world_cy = (min(ys) + max(ys)) / 2
        w = max(self._map_canvas.winfo_width(), 100)
        h = max(self._map_canvas.winfo_height(), 100)
        span_x = max(max(xs) - min(xs), 0.001)
        span_y = max(max(ys) - min(ys), 0.001)
        self._map_scale = min(w / span_x, h / span_y) * 0.85
        self._map_pan_x = 0.0
        self._map_pan_y = 0.0
        self._map_fitted = True

    def _map_reset_view(self):
        self._map_fit_all()
        self._map_redraw()

    def _map_w2c(self, wx, wy):
        """World coords → canvas pixel coords (Y flipped so +y is up)."""
        w = max(self._map_canvas.winfo_width(), 1)
        h = max(self._map_canvas.winfo_height(), 1)
        cx = (wx - self._map_world_cx) * self._map_scale + w / 2 + self._map_pan_x
        cy = -(wy - self._map_world_cy) * self._map_scale + h / 2 + self._map_pan_y
        return cx, cy

    def _map_redraw(self):
        c = self._map_canvas
        c.delete("all")
        if not self._map_points:
            c.create_text(
                max(c.winfo_width(), 200) // 2,
                max(c.winfo_height(), 200) // 2,
                text="Load a .smap or .json file to view the map",
                fill=TEXT_LO, font=FONT_LABEL)
            return

        show_ap  = self._map_show_ap.get()
        show_lm  = self._map_show_lm.get()
        show_lbl = self._map_show_labels.get()
        r_base   = max(4, min(10, int(self._map_scale * 0.6)))

        for pt in self._map_points:
            if pt["type"] == "AP" and not show_ap:
                continue
            if pt["type"] == "LM" and not show_lm:
                continue

            cx, cy = self._map_w2c(pt["x"], pt["y"])
            name       = pt["name"]
            excluded   = name in self._map_excluded
            loading    = name in self._map_loading_aps
            unloading  = name in self._map_unloading_aps
            cross_aisle = name in self._map_cross_aisle_lms

            if pt["type"] == "AP":
                r = r_base
                if cross_aisle:
                    fill, outline = "#06B6D4", "#FFFFFF"
                elif loading:
                    fill, outline = SUCCESS, "#FFFFFF"
                elif unloading:
                    fill, outline = WARNING, "#FFFFFF"
                elif excluded:
                    fill, outline = ERROR, "#FFFFFF"
                else:
                    fill, outline = ACCENT, ACCENT_ALT
                c.create_oval(cx - r, cy - r, cx + r, cy + r,
                              fill=fill, outline=outline, width=1,
                              tags=("pt", name))
                if excluded:
                    c.create_line(cx - r + 2, cy - r + 2,
                                  cx + r - 2, cy + r - 2,
                                  fill="#FFFFFF", width=1)
                    c.create_line(cx + r - 2, cy - r + 2,
                                  cx - r + 2, cy + r - 2,
                                  fill="#FFFFFF", width=1)
                elif cross_aisle or loading or unloading:
                    marker = "+" if cross_aisle else ("L" if loading else "U")
                    c.create_text(cx, cy, text=marker, fill="#FFFFFF",
                                  font=("Segoe UI", max(6, r - 1), "bold"),
                                  tags=("pt", name))
            else:
                r = max(3, r_base - 1)
                if cross_aisle:
                    fill, outline = "#06B6D4", "#FFFFFF"
                elif excluded:
                    fill, outline = ERROR, "#FFFFFF"
                else:
                    fill, outline = SUCCESS, "#34D399"
                c.create_polygon(
                    cx, cy - r, cx + r, cy,
                    cx, cy + r, cx - r, cy,
                    fill=fill, outline=outline, width=1,
                    tags=("pt", name))
                if cross_aisle:
                    c.create_text(cx, cy, text="+", fill="#FFFFFF",
                                  font=("Segoe UI", max(6, r - 1), "bold"),
                                  tags=("pt", name))

            if show_lbl:
                c.create_text(cx, cy + r_base + 7,
                              text=pt["name"], fill=TEXT_LO,
                              font=("Segoe UI", 7),
                              tags=("lbl", pt["name"]))

    def _map_press(self, event):
        self._map_sel_start  = (event.x, event.y)
        self._map_sel_active = False

    def _map_drag_select(self, event):
        if self._map_sel_start is None:
            return
        x0, y0 = self._map_sel_start
        if not self._map_sel_active:
            if math.hypot(event.x - x0, event.y - y0) > 5:
                self._map_sel_active = True
        if self._map_sel_active:
            self._map_canvas.delete("sel_rect")
            self._map_canvas.create_rectangle(
                x0, y0, event.x, event.y,
                outline=ACCENT, fill="", dash=(4, 3),
                width=2, tags="sel_rect")

    def _map_release(self, event):
        if self._map_sel_start is None:
            return
        self._map_canvas.delete("sel_rect")
        if self._map_sel_active:
            x0, y0 = self._map_sel_start
            rx0, rx1 = min(x0, event.x), max(x0, event.x)
            ry0, ry1 = min(y0, event.y), max(y0, event.y)
            changed = False
            for pt in self._map_points:
                if pt["type"] == "AP" and not self._map_show_ap.get():
                    continue
                if pt["type"] == "LM" and not self._map_show_lm.get():
                    continue
                cx, cy = self._map_w2c(pt["x"], pt["y"])
                if rx0 <= cx <= rx1 and ry0 <= cy <= ry1:
                    self._map_apply_tag(pt["name"], pt["type"])
                    changed = True
            if changed:
                self._map_update_badges()
                self._map_redraw()
        else:
            self._map_click(event)
        self._map_sel_start  = None
        self._map_sel_active = False

    def _map_click(self, event):
        threshold = max(10, int(self._map_scale * 0.8))
        best, best_d = None, threshold
        for pt in self._map_points:
            if pt["type"] == "AP" and not self._map_show_ap.get():
                continue
            if pt["type"] == "LM" and not self._map_show_lm.get():
                continue
            cx, cy = self._map_w2c(pt["x"], pt["y"])
            d = math.hypot(event.x - cx, event.y - cy)
            if d < best_d:
                best_d, best = d, pt["name"]
        if best:
            pt_type = next((p["type"] for p in self._map_points
                            if p["name"] == best), "AP")
            self._map_apply_tag(best, pt_type)
            self._map_update_badges()
            self._map_redraw()

    def _map_apply_tag(self, name: str, pt_type: str = "AP"):
        """Apply the current tag mode to a point.
        Loading/Unloading are restricted to APs only."""
        mode = self._map_tag_mode.get()
        if mode in ("loading", "unloading") and pt_type != "AP":
            return
        # Clear all existing tags for this point
        self._map_excluded.discard(name)
        self._map_loading_aps.discard(name)
        self._map_unloading_aps.discard(name)
        self._map_cross_aisle_lms.discard(name)
        if mode == "loading":
            self._map_loading_aps.add(name)
        elif mode == "unloading":
            self._map_unloading_aps.add(name)
        elif mode == "cross_aisle":
            self._map_cross_aisle_lms.add(name)
        else:
            self._map_excluded.add(name)

    def _map_update_badges(self):
        self._map_load_count_label.configure(
            text=str(len(self._map_loading_aps)))
        self._map_unload_count_label.configure(
            text=str(len(self._map_unloading_aps)))
        self._map_excl_count_label.configure(
            text=str(len(self._map_excluded)))
        self._map_ca_count_label.configure(
            text=str(len(self._map_cross_aisle_lms)))

    def _map_set_tag_mode(self, mode: str):
        self._map_tag_mode.set(mode)
        colors = {
            "loading":     SUCCESS,
            "unloading":   WARNING,
            "cross_aisle": "#06B6D4",
            "remove":      ERROR,
        }
        for val, btn in self._map_tag_btns.items():
            active = (val == mode)
            btn.configure(
                fg_color=colors[val] if active else BG_PANEL,
                text_color=TEXT_HI if active else colors[val])

    def _map_pan_start(self, event):
        self._map_drag_last = (event.x, event.y)
        self._map_canvas.configure(cursor="fleur")

    def _map_pan_move(self, event):
        if self._map_drag_last:
            dx = event.x - self._map_drag_last[0]
            dy = event.y - self._map_drag_last[1]
            self._map_pan_x += dx
            self._map_pan_y += dy
            self._map_drag_last = (event.x, event.y)
            self._map_redraw()

    def _map_pan_end(self, _):
        self._map_drag_last = None
        self._map_canvas.configure(cursor="crosshair")

    def _map_zoom(self, event):
        factor = 1.12 if event.delta > 0 else 1 / 1.12
        w = max(self._map_canvas.winfo_width(), 1)
        h = max(self._map_canvas.winfo_height(), 1)
        mx_world = (event.x - w / 2 - self._map_pan_x) / self._map_scale + self._map_world_cx
        my_world = -(event.y - h / 2 - self._map_pan_y) / self._map_scale + self._map_world_cy
        self._map_scale *= factor
        new_cx = event.x - w / 2 - (mx_world - self._map_world_cx) * self._map_scale
        new_cy = event.y - h / 2 + (my_world - self._map_world_cy) * self._map_scale
        self._map_pan_x = new_cx
        self._map_pan_y = new_cy
        self._map_redraw()
        return "break"

    def _map_on_resize(self, _):
        if not hasattr(self, "_map_points"):
            return
        if self._map_points and not self._map_fitted:
            self._map_fit_all()
        self._map_redraw()

    def _map_clear_exclusions(self):
        self._map_excluded.clear()
        self._map_loading_aps.clear()
        self._map_unloading_aps.clear()
        self._map_cross_aisle_lms.clear()
        self._map_update_badges()
        self._map_redraw()

    def _map_apply_exclusions(self):
        self.excluded_aps = list(self._map_excluded)
        count = len(self.excluded_aps)
        messagebox.showinfo(
            "Applied",
            f"{count} exclusion(s) applied.\n\n"
            "Go to Pick Sequence → Load & Analyze to regenerate with these exclusions.")

    # ── Tab3 logic ──
    def browse_file(self, entry, filetype):
        ext = filetype.replace("*", "")
        filename = filedialog.askopenfilename(
            title=f"Select {ext.upper()}", filetypes=[(f"{ext.upper()} File", filetype)])
        if filename:
            entry.delete(0, tk.END)
            entry.insert(0, filename)
            if ext.lower() == ".json":
                self.show_json_format_info()
            if ext.lower() == ".xlsx":
                self.show_excel_format_info()
                if self.create_sequence_var.get():
                    self.load_subzones_from_excel()

    def toggle_sequence_options(self):
        if self.create_sequence_var.get():
            self.sequence_options_frame.pack(fill="x", pady=(0, 8))
            self.load_subzones_from_excel()
        else:
            self.sequence_options_frame.pack_forget()

    def load_subzones_from_excel(self):
        excel_path = self._tab3_excel_picker.get()
        if not excel_path or not os.path.exists(excel_path):
            return
        try:
            df = pd.read_excel(excel_path)
            if 'subzone name' not in df.columns:
                return
            for item in self.subzone_tree.get_children():
                self.subzone_tree.delete(item)
            subzones = df['subzone name'].unique().tolist()
            for subzone in subzones:
                self.subzone_tree.insert("", "end", values=(subzone, "Left to Right"))
            self.log_tab3(f"📊 Loaded {len(subzones)} subzones for sequence configuration")
        except Exception as e:
            self.log_tab3(f"⚠️ Could not load subzones: {e}")

    def on_subzone_direction_click(self, event):
        item = self.subzone_tree.identify_row(event.y)
        col = self.subzone_tree.identify_column(event.x)
        if col == "#2":
            current = self.subzone_tree.item(item, "values")[1]
            directions = ["Left to Right", "Right to Left", "Top to Bottom", "Bottom to Top"]
            current_index = directions.index(current) if current in directions else 0
            next_index = (current_index + 1) % len(directions)
            self.subzone_tree.set(item, "Direction", directions[next_index])

    def set_all_subzone_directions(self, direction_type):
        direction_map = {
            "left_to_right": "Left to Right",
            "right_to_left": "Right to Left",
            "top_to_bottom": "Top to Bottom",
            "bottom_to_top": "Bottom to Top"
        }
        direction_text = direction_map.get(direction_type, "Left to Right")
        for item in self.subzone_tree.get_children():
            self.subzone_tree.set(item, "Direction", direction_text)

    def toggle_subzone_directions(self):
        for item in self.subzone_tree.selection():
            current = self.subzone_tree.item(item, "values")[1]
            directions = ["Left to Right", "Right to Left", "Top to Bottom", "Bottom to Top"]
            current_index = directions.index(current) if current in directions else 0
            self.subzone_tree.set(item, "Direction",
                                  directions[(current_index + 1) % len(directions)])

    def run_zone_configuration(self):
        json_path = self._tab3_json_picker.get()
        excel_path = self._tab3_excel_picker.get()
        output_file = self.tab3_output_entry.get()
        if not json_path or not excel_path:
            messagebox.showerror("Error", "Select both JSON & Excel files")
            return
        try:
            self.update_status("Loading files...")
            self.log_tab3("🔍 Loading JSON...")
            mapping = self._safe_load_json(json_path)
            if mapping is None:
                return
            self.log_tab3("🔍 Loading Excel...")
            try:
                df = pd.read_excel(excel_path)
            except Exception as e:
                messagebox.showerror("Excel Error", f"Cannot read Excel: {e}")
                return
            # ── N-Deep branch ────────────────────────────────────────
            if self.ndeep_enabled_var and self.ndeep_enabled_var.get():
                ok, msg = self.validate_excel_for_ndeep(df)
                if not ok:
                    messagebox.showerror("Invalid Excel (N-Deep)", msg)
                    self.log_tab3("❌ " + msg)
                    return
                drop_seq        = self.ndeep_drop_seq_var.get()     # "f2l" | "l2f"
                zone_scope      = self.ndeep_zone_scope_var.get()  # "BOTH"|"ENTRY"|"EXIT"
                sub_scope       = self.ndeep_sub_scope_var.get()   # "BOTH"|"ENTRY"|"EXIT"
                loc_scope       = self.ndeep_loc_scope_var.get()   # "BOTH"|"ENTRY"|"EXIT"
                entry_exit_type = self.location_entry_type.get()   # "LM" | "AP"
                self.log_tab3(f"⚙️ N-Deep mode — drop_seq={drop_seq}, "
                              f"zone={zone_scope}, subzone={sub_scope}, "
                              f"location={loc_scope}, entry/exit type={entry_exit_type}")
                self.log_tab3("🔗 Building graph…")
                graph, node_positions = self.build_graph_from_json(mapping)
                self.log_tab3(f"✅ Graph: {len(graph)} nodes, {len(node_positions)} positions")
                self.log_tab3("🔄 Generating N-Deep configuration…")
                result = self.generate_ndeep_zone_configuration(
                    mapping, df, drop_seq,
                    zone_scope, sub_scope, loc_scope,
                    entry_exit_type, graph, node_positions)
                self.log_tab3("💾 Saving JSON…")
                with open(output_file, "w", encoding="utf-8") as f:
                    json.dump(result, f, indent=2)
                messagebox.showinfo("Success",
                                    f"N-Deep zone config saved to {output_file}")
                self.update_status("N-Deep zone config generated")
                return  # done — skip the standard flow below

            # ── Standard (non-N-Deep) flow ───────────────────────────
            ok, msg = self.validate_excel_for_zone(df)
            if not ok:
                messagebox.showerror("Invalid Excel", msg)
                self.log_tab3("❌ " + msg)
                return
            entry_exit_type = self.location_entry_type.get()
            self.log_tab3(f"📍 Using {entry_exit_type}s for location entry/exit points")
            self.log_tab3("🔗 Building graph for path distance calculations...")
            graph, node_positions = self.build_graph_from_json(mapping)
            self.log_tab3(f"✅ Graph built with {len(graph)} nodes and {len(node_positions)} positions")
            if self.create_sequence_var.get():
                self.log_tab3("🔄 User requested sequences for subzones")
                sequence_configs = {}
                for item in self.subzone_tree.get_children():
                    subzone, direction = self.subzone_tree.item(item, "values")
                    direction_map = {
                        "Left to Right": "left_to_right",
                        "Right to Left": "right_to_left",
                        "Top to Bottom": "top_to_bottom",
                        "Bottom to Top": "bottom_to_top"
                    }
                    sequence_configs[subzone] = {"direction": direction_map.get(direction, "left_to_right")}
                    self.log_tab3(f"⚙️ Subzone '{subzone}': {direction}")
                self.log_tab3("🔄 Generating configuration with sequences...")
                result = self.generate_zone_configuration_with_sequences(
                    mapping, df, sequence_configs, entry_exit_type, graph, node_positions)
                self.log_tab3(f"✅ Generated configuration with sequences for {len(sequence_configs)} subzones")
            else:
                self.log_tab3("🔄 Generating configuration without sequences...")
                result = self.generate_zone_configuration(mapping, df, entry_exit_type, graph, node_positions)
            self.log_tab3("💾 Saving JSON...")
            with open(output_file, "w", encoding="utf-8") as f:
                json.dump(result, f, indent=2)
            messagebox.showinfo("Success", f"Saved to {output_file}")
            self.update_status("Zone config generated")
        except Exception as e:
            self.log_tab3(str(e))
            messagebox.showerror("Error", str(e))
            self.update_status("Error generating zone config")

    def generate_zone_configuration_with_sequences(self, mapping_data, df, sequence_configs,
                                                   entry_exit_type="LM", graph=None, node_positions=None):
        aps, lms = [], []
        if "advancedPointList" in mapping_data:
            for point in mapping_data["advancedPointList"]:
                if point.get("className") == "ActionPoint":
                    aps.append(point)
                elif point.get("className") == "LocationMark":
                    lms.append(point)
        data = {"LocationDataMapping": {"zones": []}}
        zones = []
        current_zone, subzones = None, []
        grouped = df.groupby(['Zone name', 'subzone name'])
        for (zone_name, subzone_name), group in grouped:
            row = group.iloc[0]
            zone_entry = str(row['zone entry'])
            zone_exit = str(row['zone exit'])
            sub_entry = str(row['subzone entery'])
            sub_exit = str(row['subzone exit'])
            min_x, max_x = group['min x'].min(), group['max x'].max()
            min_y, max_y = group['min y'].min(), group['max y'].max()
            if current_zone is None or current_zone["name"] != zone_name:
                if current_zone is not None:
                    current_zone["subZones"] = subzones
                    zones.append(current_zone)
                current_zone = {
                    "name": zone_name,
                    "entryPoint": zone_entry,
                    "exitPoint": zone_exit,
                    "locationSelectionStrategy": "SEQUENTIAL",
                    "subZones": []
                }
                subzones = []
            subzone_aps = []
            for ap in aps:
                if min_x <= ap["pos"]["x"] <= max_x and min_y <= ap["pos"]["y"] <= max_y:
                    subzone_aps.append({"name": ap["instanceName"], "x": ap["pos"]["x"], "y": ap["pos"]["y"]})
            if subzone_name in sequence_configs:
                config = sequence_configs[subzone_name]
                direction = config["direction"]
                if direction in ["left_to_right", "right_to_left"]:
                    subzone_aps.sort(key=lambda ap: (round(ap["y"], 6), ap["x"]))
                    if direction == "right_to_left":
                        grouped_by_y = {}
                        for ap in subzone_aps:
                            y_key = round(ap["y"], 6)
                            grouped_by_y.setdefault(y_key, []).append(ap)
                        sorted_aps = []
                        for y_key in sorted(grouped_by_y.keys()):
                            row_aps = grouped_by_y[y_key]
                            sorted_aps.extend(sorted(row_aps, key=lambda ap: ap["x"], reverse=True))
                        subzone_aps = sorted_aps
                else:
                    subzone_aps.sort(key=lambda ap: (round(ap["x"], 6), ap["y"]))
                    if direction == "bottom_to_top":
                        grouped_by_x = {}
                        for ap in subzone_aps:
                            x_key = round(ap["x"], 6)
                            grouped_by_x.setdefault(x_key, []).append(ap)
                        sorted_aps = []
                        for x_key in sorted(grouped_by_x.keys()):
                            col_aps = grouped_by_x[x_key]
                            sorted_aps.extend(sorted(col_aps, key=lambda ap: ap["y"], reverse=True))
                        subzone_aps = sorted_aps
            locations = []
            for idx, ap in enumerate(subzone_aps, 1):
                if entry_exit_type == "LM":
                    entry_point = self.find_closest_point_by_path_distance(
                        ap["name"], lms, graph, node_positions, "LM")
                    exit_point = entry_point
                else:
                    entry_point = self.find_closest_point_by_path_distance(
                        ap["name"], aps, graph, node_positions, "AP")
                    exit_point = entry_point
                locations.append({
                    "customerLocation": f"{subzone_name}-{idx:03d}",
                    "fleetManagerLocation": ap['name'],
                    "locationType": "BOTH",
                    "entryPoint": entry_point,
                    "exitPoint": exit_point
                })
            subzones.append({
                "name": subzone_name,
                "entryPoint": sub_entry,
                "exitPoint": sub_exit,
                "locationSelectionStrategy": "SEQUENTIAL",
                "locations": locations
            })
        if current_zone:
            current_zone["subZones"] = subzones
            zones.append(current_zone)
        data["LocationDataMapping"]["zones"] = zones
        return data

    def generate_zone_configuration(self, mapping_data, df, entry_exit_type="LM", graph=None, node_positions=None):
        aps, lms = [], []
        if "advancedPointList" in mapping_data:
            for point in mapping_data["advancedPointList"]:
                if point.get("className") == "ActionPoint":
                    aps.append(point)
                elif point.get("className") == "LocationMark":
                    lms.append(point)
        data = {"LocationDataMapping": {"zones": []}}
        zones = []
        current_zone, subzones = None, []
        for _, row in df.iterrows():
            zone = str(row['Zone name'])
            zone_entry = str(row['zone entry'])
            zone_exit = str(row['zone exit'])
            subzone = str(row['subzone name'])
            sub_entry = str(row['subzone entery'])
            sub_exit = str(row['subzone exit'])
            min_x, max_x = float(row['min x']), float(row['max x'])
            min_y, max_y = float(row['min y']), float(row['max y'])
            if current_zone is None or current_zone["name"] != zone:
                if current_zone is not None:
                    current_zone["subZones"] = subzones
                    zones.append(current_zone)
                current_zone = {
                    "name": zone,
                    "entryPoint": zone_entry,
                    "exitPoint": zone_exit,
                    "locationSelectionStrategy": "RANDOM",
                    "subZones": []
                }
                subzones = []
            locations = []
            idx = 1
            for ap in aps:
                if min_x <= ap["pos"]["x"] <= max_x and min_y <= ap["pos"]["y"] <= max_y:
                    if entry_exit_type == "LM":
                        entry_point = self.find_closest_point_by_path_distance(
                            ap["instanceName"], lms, graph, node_positions, "LM")
                        exit_point = entry_point
                    else:
                        entry_point = self.find_closest_point_by_path_distance(
                            ap["instanceName"], aps, graph, node_positions, "AP")
                        exit_point = entry_point
                    locations.append({
                        "customerLocation": f"{subzone}-{idx}",
                        "fleetManagerLocation": ap['instanceName'],
                        "locationType": "BOTH",
                        "entryPoint": entry_point,
                        "exitPoint": exit_point
                    })
                    idx += 1
            subzones.append({
                "name": subzone,
                "entryPoint": sub_entry,
                "exitPoint": sub_exit,
                "locationSelectionStrategy": "RANDOM",
                "locations": locations
            })
        if current_zone:
            current_zone["subZones"] = subzones
            zones.append(current_zone)
        data["LocationDataMapping"]["zones"] = zones
        return data

    def find_closest_point_by_path_distance(self, source_name, target_points, graph, node_positions, point_type="LM"):
        if not target_points:
            return f"{point_type}1"
        source_pos = None
        if node_positions and source_name in node_positions:
            source_pos = node_positions[source_name]
        else:
            for point in target_points:
                if point.get("instanceName") == source_name:
                    pos = point.get("pos", {})
                    if "x" in pos and "y" in pos:
                        source_pos = (float(pos["x"]), float(pos["y"]))
                        break
        if source_pos is None:
            return f"{point_type}1"
        distances, _ = self.dijkstra(source_name, graph)
        best_point = None
        best_distance = float('inf')
        for point in target_points:
            target_name = point.get("instanceName")
            if target_name == source_name:
                continue
            distance = distances.get(target_name, float('inf'))
            if distance == float('inf') or math.isinf(distance):
                pos = point.get("pos", {})
                if "x" not in pos or "y" not in pos:
                    continue
                target_pos = (float(pos["x"]), float(pos["y"]))
                distance = math.sqrt((source_pos[0] - target_pos[0])**2 +
                                     (source_pos[1] - target_pos[1])**2)
            if distance < best_distance:
                best_distance = distance
                best_point = target_name
        if best_point is None:
            for point in target_points:
                target_name = point.get("instanceName")
                if target_name == source_name:
                    continue
                pos = point.get("pos", {})
                if "x" not in pos or "y" not in pos:
                    continue
                target_pos = (float(pos["x"]), float(pos["y"]))
                distance = math.sqrt((source_pos[0] - target_pos[0])**2 +
                                     (source_pos[1] - target_pos[1])**2)
                if distance < best_distance:
                    best_distance = distance
                    best_point = target_name
        return best_point or f"{point_type}1"


# ─────────────────────────────────────────────
# Shim: makes new picker widgets look like old tk.Entry to original logic
# ─────────────────────────────────────────────

class _EntryShim:
    def __init__(self, picker):
        self._picker = picker

    def get(self):
        return self._picker.get()

    def delete(self, *args):
        self._picker.entry.delete(*args)

    def insert(self, pos, val):
        self._picker.entry.insert(pos, val)


# ─────────────────────────────────────────────
# Entry point
# ─────────────────────────────────────────────

if __name__ == "__main__":
    root = ctk.CTk()
    app = WarehouseToolsApp(root)
    root.mainloop()
