"""
warehouse_logic.py — Pure Python business logic for Warehouse Tools.
No tkinter, no UI imports. All functions are stateless (graph/data passed in).
"""

import json
import math
import os
import random
import time
from collections import defaultdict, deque

import heapq
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment


# ─────────────────────────────────────────────
# Validation helpers
# ─────────────────────────────────────────────

def validate_json_for_markers(data):
    if not isinstance(data, dict):
        return False, "Markers JSON must be an object"
    if "markers" not in data:
        return False, "Missing 'markers' array in JSON"
    markers = data.get("markers", [])
    if not isinstance(markers, list) or not markers:
        return False, "'markers' must be a non-empty list"
    for m in markers:
        if not isinstance(m, dict):
            continue
        infos = m.get("markInfos", [])
        if isinstance(infos, list):
            for info in infos:
                if isinstance(info, dict) and "x" in info and "y" in info:
                    return True, "OK"
    return False, "No valid markInfos with x/y found in markers"


def validate_json_for_pick_sequence(data):
    if not isinstance(data, dict):
        return False, "JSON root must be an object"
    if "advancedPointList" not in data:
        return False, "Missing 'advancedPointList'"
    if "advancedCurveList" not in data:
        return False, "Missing 'advancedCurveList'"
    ap_found = False
    for p in data.get("advancedPointList", []):
        if p.get("className") == "ActionPoint":
            pos = p.get("pos") or {}
            if "x" in pos and "y" in pos:
                ap_found = True
                break
    if not ap_found:
        return False, "No ActionPoint with valid pos found in 'advancedPointList'"
    for c in data.get("advancedCurveList", []):
        s = c.get("startPos", {})
        e = c.get("endPos", {})
        if not s.get("instanceName") or not e.get("instanceName"):
            return False, "A curve entry is missing start/end instanceName"
        if "pos" not in s or "pos" not in e:
            return False, "A curve entry is missing position ('pos') for start or end"
        if "x" not in s.get("pos", {}) or "y" not in s.get("pos", {}):
            return False, "Start position missing x/y in a curve"
        if "x" not in e.get("pos", {}) or "y" not in e.get("pos", {}):
            return False, "End position missing x/y in a curve"
    return True, "OK"


def validate_excel_for_zone(df):
    required = ['Zone name', 'zone entry', 'zone exit', 'subzone name',
                'subzone entery', 'subzone exit', 'min x', 'max x', 'min y', 'max y']
    missing = [c for c in required if c not in df.columns]
    if missing:
        return False, f"Missing columns: {missing}"
    try:
        for col in ['min x', 'max x', 'min y', 'max y']:
            pd.to_numeric(df[col], errors='raise')
    except Exception as e:
        return False, f"Coordinate columns must be numeric: {e}"
    return True, "OK"


def validate_excel_for_ndeep(df):
    required = ['Zone name', 'zone entry', 'zone exit',
                'min x', 'max x', 'min y', 'max y']
    missing = [c for c in required if c not in df.columns]
    if missing:
        return False, f"Missing columns for N-Deep: {missing}"
    try:
        for col in ['min x', 'max x', 'min y', 'max y']:
            pd.to_numeric(df[col], errors='raise')
    except Exception as e:
        return False, f"Coordinate columns must be numeric: {e}"
    return True, "OK"


# ─────────────────────────────────────────────
# Graph construction
# ─────────────────────────────────────────────

def build_graph_from_json(json_data):
    graph = {}
    node_positions = {}
    for p in json_data.get("advancedPointList", []):
        inst = p.get("instanceName")
        pos = p.get("pos", {})
        if inst and "x" in pos and "y" in pos:
            node_positions[inst] = (float(pos["x"]), float(pos["y"]))
            graph.setdefault(inst, {})
    for curve in json_data.get("advancedCurveList", []):
        try:
            start = curve.get("startPos", {}).get("instanceName")
            end = curve.get("endPos", {}).get("instanceName")
            sx = curve.get("startPos", {}).get("pos", {}).get("x")
            sy = curve.get("startPos", {}).get("pos", {}).get("y")
            ex = curve.get("endPos", {}).get("pos", {}).get("x")
            ey = curve.get("endPos", {}).get("pos", {}).get("y")
            if start and start not in node_positions and sx is not None and sy is not None:
                node_positions[start] = (float(sx), float(sy))
                graph.setdefault(start, {})
            if end and end not in node_positions and ex is not None and ey is not None:
                node_positions[end] = (float(ex), float(ey))
                graph.setdefault(end, {})
        except Exception:
            continue
    for curve in json_data.get("advancedCurveList", []):
        start = curve.get("startPos", {}).get("instanceName")
        end = curve.get("endPos", {}).get("instanceName")
        if not start or not end:
            continue
        s_pos = curve.get("startPos", {}).get("pos", {})
        e_pos = curve.get("endPos", {}).get("pos", {})
        if "x" not in s_pos or "y" not in s_pos or "x" not in e_pos or "y" not in e_pos:
            continue
        sx, sy = float(s_pos["x"]), float(s_pos["y"])
        ex, ey = float(e_pos["x"]), float(e_pos["y"])
        weight = math.hypot(ex - sx, ey - sy)
        if weight <= 0:
            continue
        graph.setdefault(start, {})[end] = min(weight, graph.get(start, {}).get(end, weight))
    return graph, node_positions


# ─────────────────────────────────────────────
# Shortest path
# ─────────────────────────────────────────────

def dijkstra(source, graph):
    if source not in graph:
        return {}, {}
    dist = {n: float('inf') for n in graph}
    prev = {n: None for n in graph}
    dist[source] = 0.0
    heap = [(0.0, source)]
    while heap:
        d, u = heapq.heappop(heap)
        if d > dist[u]:
            continue
        for v, w in graph.get(u, {}).items():
            nd = d + w
            if nd < dist.get(v, float('inf')):
                dist[v] = nd
                prev[v] = u
                heapq.heappush(heap, (nd, v))
    return dist, prev


def calculate_distance(a_name, b_name, graph, node_positions):
    if a_name == b_name:
        return 0.0
    if a_name not in graph or b_name not in graph:
        pos_a = node_positions.get(a_name)
        pos_b = node_positions.get(b_name)
        if pos_a and pos_b:
            return math.hypot(pos_a[0] - pos_b[0], pos_a[1] - pos_b[1])
        return float('inf')
    dist, _ = dijkstra(a_name, graph)
    return dist.get(b_name, float('inf'))


# ─────────────────────────────────────────────
# Tab 1 — Each Pick (Marker Extraction)
# ─────────────────────────────────────────────

def extract_markers(data, marker_type):
    return [
        {"code": m.get("code"), "x": info.get("x"), "y": info.get("y")}
        for m in data.get("markers", []) if m.get("type") == marker_type
        for info in m.get("markInfos", [])
        if info.get("x") is not None and info.get("y") is not None
    ]


def process_markers(markers):
    df = pd.DataFrame(markers)
    if df.empty:
        return pd.DataFrame(), set()
    grouped = {}
    for x, group in df.groupby("x"):
        sorted_group = group.sort_values("y")
        grouped[x] = sorted_group["code"].tolist()
    max_len = max(len(codes) for codes in grouped.values())
    aligned = {x: codes + [None] * (max_len - len(codes)) for x, codes in grouped.items()}
    result_df = pd.DataFrame(aligned)
    all_codes = set(df["code"])
    used_codes = set(code for lst in grouped.values() for code in lst if code)
    missing_codes = all_codes - used_codes
    return result_df, missing_codes


def run_marker_extraction(json_data, include_work=True, include_resource=True):
    """
    Returns: {"sheets": [(name, df), ...], "log": [str, ...], "error": str|None}
    """
    log = []
    ok, msg = validate_json_for_markers(json_data)
    if not ok:
        return {"error": msg, "log": log}

    log.append("JSON structure validated")
    work_sheets = []
    validation_data = {}

    if include_work:
        log.append("Processing WorkMarkers...")
        wm = extract_markers(json_data, "WorkMarker")
        result, missing = process_markers(wm)
        work_sheets.append(("WorkMarkers_Grouped", result))
        if missing:
            validation_data["Missing_WorkMarkers"] = list(missing)
        log.append(f"Found {len(wm)} WorkMarkers")

    if include_resource:
        log.append("Processing ResourceMarkers...")
        rm = extract_markers(json_data, "ResourceMarker")
        result, missing = process_markers(rm)
        work_sheets.append(("ResourceMarkers_Grouped", result))
        if missing:
            validation_data["Missing_ResourceMarkers"] = list(missing)
        log.append(f"Found {len(rm)} ResourceMarkers")

    return {"sheets": work_sheets, "validation": validation_data, "log": log, "error": None}


def save_marker_extraction_to_excel(sheets, validation_data, output_path):
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets:
            if isinstance(df, pd.DataFrame) and df.empty:
                pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)
            else:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        for sheet, missing_codes in validation_data.items():
            pd.DataFrame({"Missing Codes": missing_codes}).to_excel(
                writer, sheet_name=sheet, index=False)


# ─────────────────────────────────────────────
# Tab 2 — Case Pick (Pick Sequence)
# ─────────────────────────────────────────────

def get_action_points(json_data):
    aps = []
    for ap in json_data.get("advancedPointList", []):
        if ap.get("className") == "ActionPoint":
            pos = ap.get("pos", {})
            if "x" in pos and "y" in pos:
                aps.append({"name": ap.get("instanceName"),
                            "x": float(pos["x"]), "y": float(pos["y"])})
    return aps


def auto_group_aisles(aps, orientation="horizontal", gap=0.2):
    if not aps:
        return [], "x"
    if orientation == "horizontal":
        aps.sort(key=lambda ap: (round(ap["y"], 6), ap["x"]))
        axis = "y"
        sort_axis = "x"
    else:
        aps.sort(key=lambda ap: (round(ap["x"], 6), ap["y"]))
        axis = "x"
        sort_axis = "y"
    aisles = []
    current = [aps[0]]
    last = aps[0][axis]
    for ap in aps[1:]:
        if abs(ap[axis] - last) > gap:
            aisles.append(current)
            current = []
        current.append(ap)
        last = ap[axis]
    aisles.append(current)
    for aisle in aisles:
        aisle.sort(key=lambda ap: ap[sort_axis])
    return aisles, sort_axis


def split_aisles_by_markers(aisles, sort_axis, all_points, cross_aisle_names):
    """
    Split aisles at coordinates of manually-tagged cross-aisle markers.
    all_points: list of {"name", "x", "y"} dicts.
    cross_aisle_names: set of point names marked as cross-aisle.
    """
    cross_positions = sorted(
        pt[sort_axis] for pt in all_points
        if pt["name"] in cross_aisle_names
    )
    if not cross_positions:
        return aisles

    output = []
    for aisle in aisles:
        aisle_sorted = sorted(aisle, key=lambda ap: ap[sort_axis])
        coords = [ap[sort_axis] for ap in aisle_sorted]

        split_indices = set()
        for cp in cross_positions:
            for i in range(len(coords) - 1):
                lo = min(coords[i], coords[i + 1])
                hi = max(coords[i], coords[i + 1])
                if lo < cp < hi:
                    split_indices.add(i + 1)
                    break

        if not split_indices:
            output.append(aisle_sorted)
            continue

        start = 0
        for idx in sorted(split_indices):
            section = aisle_sorted[start:idx]
            if section:
                output.append(section)
            start = idx
        last = aisle_sorted[start:]
        if last:
            output.append(last)

    return output


def detect_cross_aisles(aisle, sort_axis, sensitivity=1.5, min_length=3):
    if len(aisle) < min_length:
        return [aisle], []
    distances = [abs(aisle[i + 1][sort_axis] - aisle[i][sort_axis])
                 for i in range(len(aisle) - 1)]
    if not distances:
        return [aisle], []
    arr = np.array(distances)
    q25, q50, q75 = np.percentile(arr, [25, 50, 75])
    iqr = q75 - q25
    threshold = (q75 + sensitivity * iqr) if iqr > 0 else q50 * 2.5
    threshold = max(threshold, q50 * 1.8)
    splits = sorted(i for i, d in enumerate(distances) if d > threshold)
    if not splits:
        return [aisle], []
    boundaries = [0] + [s + 1 for s in splits] + [len(aisle)]
    section_sizes = [boundaries[k + 1] - boundaries[k]
                     for k in range(len(boundaries) - 1)]
    validated = [
        s for s, (sz_l, sz_r) in zip(splits, zip(section_sizes, section_sizes[1:]))
        if sz_l >= 2 and sz_r >= 2
    ]
    if not validated:
        return [aisle], []
    idxs = [i + 1 for i in validated]
    subaisles, start = [], 0
    for idx in idxs:
        subaisles.append(aisle[start:idx])
        start = idx
    subaisles.append(aisle[start:])
    return subaisles, []


def detect_all_cross_aisles(aisles, sort_axis, sensitivity=1.5):
    output = []
    for aisle in aisles:
        if len(aisle) >= 3:
            subs, _ = detect_cross_aisles(aisle, sort_axis, sensitivity)
            output.extend(subs)
        else:
            output.append(aisle)
    return output


def optimize_aisle_order(aisles, loading_point_name, graph, node_positions):
    if not aisles:
        return []
    dist_from_lp, _ = dijkstra(loading_point_name, graph) if loading_point_name in graph else ({}, {})
    aisle_scores = []
    for aisle in aisles:
        min_d = float('inf')
        for ap in aisle:
            d = dist_from_lp.get(ap['name'], float('inf'))
            if math.isinf(d):
                d = calculate_distance(loading_point_name, ap['name'], graph, node_positions)
            if d < min_d:
                min_d = d
        aisle_scores.append((min_d, aisle))
    aisle_scores.sort(key=lambda x: (math.isinf(x[0]), x[0]))
    return [a for _, a in aisle_scores]


def optimize_aisle_direction(aisle, loading_point_name, prev_end_name, graph, node_positions):
    first_ap = aisle[0]
    last_ap = aisle[-1]
    ref = prev_end_name if prev_end_name else loading_point_name
    d_first = calculate_distance(ref, first_ap['name'], graph, node_positions)
    d_last = calculate_distance(ref, last_ap['name'], graph, node_positions)
    if d_first <= d_last:
        return aisle, last_ap['name']
    else:
        return list(reversed(aisle)), first_ap['name']


def generate_automatic_sequence(aisles, loading_point_name, graph, node_positions, log_fn=None):
    ordered_aisles = optimize_aisle_order(aisles, loading_point_name, graph, node_positions)
    final_sequence = []
    current_position = loading_point_name
    for aisle in ordered_aisles:
        aisle_nodes = [ap for ap in aisle if ap['name'] in node_positions]
        if not aisle_nodes:
            if log_fn:
                log_fn(f"Skipping aisle (no graph nodes): {[ap['name'] for ap in aisle]}")
            continue
        optimized_aisle, end_name = optimize_aisle_direction(
            aisle_nodes, loading_point_name, current_position, graph, node_positions)
        start_ap = optimized_aisle[0]
        dist_to_start = calculate_distance(current_position, start_ap['name'], graph, node_positions)
        if math.isinf(dist_to_start):
            if log_fn:
                log_fn(f"Skipping aisle (unreachable from current pos): {start_ap['name']}")
            continue
        final_sequence.extend(optimized_aisle)
        current_position = end_name
    return final_sequence


def find_closest_lm(ap_name, json_data, graph, node_positions):
    """Return name of the LM closest to ap_name via graph path, fallback to Euclidean."""
    lm_positions = {}
    for pt in json_data.get("advancedPointList", []):
        if pt.get("className") != "ActionPoint":
            pos = pt.get("pos", {})
            inst = pt.get("instanceName")
            if inst and "x" in pos and "y" in pos:
                lm_positions[inst] = (float(pos["x"]), float(pos["y"]))
    if not lm_positions:
        return None
    if ap_name in graph:
        dist, _ = dijkstra(ap_name, graph)
        reachable = {n: d for n, d in dist.items()
                     if n in lm_positions and d < float("inf")}
        if reachable:
            return min(reachable, key=reachable.get)
    if ap_name not in node_positions:
        return None
    ax, ay = node_positions[ap_name]
    return min(lm_positions,
               key=lambda n: math.hypot(lm_positions[n][0] - ax,
                                        lm_positions[n][1] - ay))


def save_sequence_to_excel(flat_sequence, output_path,
                            loading_aps, unloading_aps,
                            json_data, graph, node_positions,
                            duplication=0, template_path="",
                            aisle_groups=None):
    """
    flat_sequence: list of AP name strings (already in order).
    loading_aps / unloading_aps: sets of AP names.
    aisle_groups: optional list of (aisle_name, [ap_name, ...]) tuples — if
        provided, an extra "aisle_AP" sheet is written with one column per
        aisle (header = aisle name, rows = APs in sequence).
    """
    dup_count = duplication if duplication >= 2 else 1

    def _write_aisle_ap_sheet(wb):
        """Add/refresh the aisle_AP sheet — one column per aisle."""
        if not aisle_groups:
            return
        if "aisle_AP" in wb.sheetnames:
            del wb["aisle_AP"]
        ws = wb.create_sheet("aisle_AP")
        for col_idx, (name, aps) in enumerate(aisle_groups, start=1):
            ws.cell(1, col_idx, name)
            for row_idx, ap_name in enumerate(aps, start=2):
                ws.cell(row_idx, col_idx, ap_name)

    def _write_station_rows(ws):
        row = 2
        for group, label, ap_set in [
            ("LOADING",   "loading",   sorted(loading_aps)),
            ("UNLOADING", "unloading", sorted(unloading_aps)),
        ]:
            for i, ap_name in enumerate(ap_set, 1):
                wm = find_closest_lm(ap_name, json_data, graph, node_positions)
                ws.cell(row, 1, group)
                ws.cell(row, 2, None)
                ws.cell(row, 3, f"{group}_{i}")
                ws.cell(row, 4, label)
                ws.cell(row, 5, "stp")
                ws.cell(row, 6, ap_name)
                ws.cell(row, 7, wm)
                ws.cell(row, 8, 1)
                ws.cell(row, 9, True)
                ws.cell(row, 10, 1)
                ws.cell(row, 11, 8)
                ws.cell(row, 12, "SELF")
                row += 1

    if template_path and os.path.exists(template_path):
        wb = load_workbook(template_path)
        ws_dup = wb["location-marker-mapping"]
        for r in range(2, ws_dup.max_row + 1):
            for c in range(1, ws_dup.max_column + 1):
                ws_dup.cell(r, c).value = None
        row = 2
        for ap_name in flat_sequence:
            for _ in range(dup_count):
                ws_dup.cell(row, 1).value = None
                ws_dup.cell(row, 2).value = ap_name
                ws_dup.cell(row, 3).value = False
                row += 1

        ws_seq = wb["marker-sequence"]
        for r in range(2, ws_seq.max_row + 1):
            for c in range(1, ws_seq.max_column + 1):
                ws_seq.cell(r, c).value = None
        for i, ap_name in enumerate(flat_sequence, 2):
            ws_seq.cell(i, 1).value = None
            ws_seq.cell(i, 2).value = ap_name

        ws_sc = wb["station-config"]
        for r in range(2, ws_sc.max_row + 1):
            for col in range(1, ws_sc.max_column + 1):
                ws_sc.cell(r, col).value = None
        _write_station_rows(ws_sc)
        _write_aisle_ap_sheet(wb)
        wb.save(output_path)
    else:
        wb = Workbook()
        ws_dup = wb.active
        ws_dup.title = "location-marker-mapping"
        ws_dup.cell(1, 1, "Location")
        ws_dup.cell(1, 2, "Marker Code")
        ws_dup.cell(1, 3, "Is Smart")
        row = 2
        for ap_name in flat_sequence:
            for _ in range(dup_count):
                ws_dup.cell(row, 1).value = None
                ws_dup.cell(row, 2).value = ap_name
                ws_dup.cell(row, 3).value = False
                row += 1

        ws_seq = wb.create_sheet("marker-sequence")
        ws_seq.cell(1, 1, "Seq")
        ws_seq.cell(1, 2, "Marker Code")
        for i, ap_name in enumerate(flat_sequence, 2):
            ws_seq.cell(i, 2, ap_name)

        ws_aisle = wb.create_sheet("aisle-sequence")
        ws_aisle.cell(1, 1, "Seq")
        ws_aisle.cell(1, 2, "Aisle")

        ws_mc = wb.create_sheet("marker-config")
        for col, h in enumerate(
                ["MarkerId", "Name", "Type", "Queue Capacity", "Is Active"], 1):
            ws_mc.cell(1, col, h)

        ws_sc = wb.create_sheet("station-config")
        for col, h in enumerate(
                ["stationGroupUid", "prePoint", "stationUid", "labels",
                 "stationType", "stationMarker", "waitMarker", "priority",
                 "active", "capacity", "height", "Ownership"], 1):
            ws_sc.cell(1, col, h)
        _write_station_rows(ws_sc)
        _write_aisle_ap_sheet(wb)
        wb.save(output_path)


# ─────────────────────────────────────────────
# Tab 3 — Point to Point (Zone Configuration)
# ─────────────────────────────────────────────

def find_closest_point_by_path_distance(source_name, target_points, graph,
                                         node_positions, point_type="LM"):
    if not target_points:
        return f"{point_type}1"
    source_pos = node_positions.get(source_name)
    if source_pos is None:
        for point in target_points:
            if point.get("instanceName") == source_name:
                pos = point.get("pos", {})
                if "x" in pos and "y" in pos:
                    source_pos = (float(pos["x"]), float(pos["y"]))
                    break
    if source_pos is None:
        return f"{point_type}1"

    distances, _ = dijkstra(source_name, graph)
    best_point, best_distance = None, float('inf')

    for point in target_points:
        target_name = point.get("instanceName")
        if target_name == source_name:
            continue
        distance = distances.get(target_name, float('inf'))
        if math.isinf(distance):
            pos = point.get("pos", {})
            if "x" not in pos or "y" not in pos:
                continue
            target_pos = (float(pos["x"]), float(pos["y"]))
            distance = math.hypot(source_pos[0] - target_pos[0],
                                  source_pos[1] - target_pos[1])
        if distance < best_distance:
            best_distance = distance
            best_point = target_name

    if best_point is None:
        for point in target_points:
            target_name = point.get("instanceName")
            if target_name == source_name:
                continue
            pos = point.get("pos", {})
            if "x" not in pos or "y" not in pos:
                continue
            target_pos = (float(pos["x"]), float(pos["y"]))
            distance = math.hypot(source_pos[0] - target_pos[0],
                                  source_pos[1] - target_pos[1])
            if distance < best_distance:
                best_distance = distance
                best_point = target_name

    return best_point or f"{point_type}1"


def generate_zone_configuration(mapping_data, df, entry_exit_type="LM",
                                  graph=None, node_positions=None):
    aps, lms = [], []
    for point in mapping_data.get("advancedPointList", []):
        if point.get("className") == "ActionPoint":
            aps.append(point)
        elif point.get("className") == "LocationMark":
            lms.append(point)
    data = {"LocationDataMapping": {"zones": []}}
    zones = []
    current_zone, subzones = None, []
    for _, row in df.iterrows():
        zone = str(row['Zone name'])
        zone_entry = str(row['zone entry'])
        zone_exit = str(row['zone exit'])
        subzone = str(row['subzone name'])
        sub_entry = str(row['subzone entery'])
        sub_exit = str(row['subzone exit'])
        min_x, max_x = float(row['min x']), float(row['max x'])
        min_y, max_y = float(row['min y']), float(row['max y'])
        if current_zone is None or current_zone["name"] != zone:
            if current_zone is not None:
                current_zone["subZones"] = subzones
                zones.append(current_zone)
            current_zone = {
                "name": zone,
                "entryPoint": zone_entry,
                "exitPoint": zone_exit,
                "locationSelectionStrategy": "RANDOM",
                "subZones": []
            }
            subzones = []
        locations = []
        idx = 1
        for ap in aps:
            if min_x <= ap["pos"]["x"] <= max_x and min_y <= ap["pos"]["y"] <= max_y:
                if entry_exit_type == "LM":
                    ep = find_closest_point_by_path_distance(
                        ap["instanceName"], lms, graph, node_positions, "LM")
                else:
                    ep = find_closest_point_by_path_distance(
                        ap["instanceName"], aps, graph, node_positions, "AP")
                locations.append({
                    "customerLocation": f"{subzone}-{idx}",
                    "fleetManagerLocation": ap['instanceName'],
                    "locationType": "BOTH",
                    "entryPoint": ep,
                    "exitPoint": ep
                })
                idx += 1
        subzones.append({
            "name": subzone,
            "entryPoint": sub_entry,
            "exitPoint": sub_exit,
            "locationSelectionStrategy": "RANDOM",
            "locations": locations
        })
    if current_zone:
        current_zone["subZones"] = subzones
        zones.append(current_zone)
    data["LocationDataMapping"]["zones"] = zones
    return data


def generate_zone_configuration_with_sequences(mapping_data, df, sequence_configs,
                                                entry_exit_type="LM",
                                                graph=None, node_positions=None):
    aps, lms = [], []
    for point in mapping_data.get("advancedPointList", []):
        if point.get("className") == "ActionPoint":
            aps.append(point)
        elif point.get("className") == "LocationMark":
            lms.append(point)
    data = {"LocationDataMapping": {"zones": []}}
    zones = []
    current_zone, subzones = None, []
    grouped = df.groupby(['Zone name', 'subzone name'])
    for (zone_name, subzone_name), group in grouped:
        row = group.iloc[0]
        zone_entry = str(row['zone entry'])
        zone_exit = str(row['zone exit'])
        sub_entry = str(row['subzone entery'])
        sub_exit = str(row['subzone exit'])
        min_x, max_x = group['min x'].min(), group['max x'].max()
        min_y, max_y = group['min y'].min(), group['max y'].max()
        if current_zone is None or current_zone["name"] != zone_name:
            if current_zone is not None:
                current_zone["subZones"] = subzones
                zones.append(current_zone)
            current_zone = {
                "name": zone_name,
                "entryPoint": zone_entry,
                "exitPoint": zone_exit,
                "locationSelectionStrategy": "SEQUENTIAL",
                "subZones": []
            }
            subzones = []
        subzone_aps = []
        for ap in aps:
            if min_x <= ap["pos"]["x"] <= max_x and min_y <= ap["pos"]["y"] <= max_y:
                subzone_aps.append({
                    "name": ap["instanceName"],
                    "x": ap["pos"]["x"],
                    "y": ap["pos"]["y"]
                })
        if subzone_name in sequence_configs:
            direction = sequence_configs[subzone_name]["direction"]
            if direction in ["left_to_right", "right_to_left"]:
                subzone_aps.sort(key=lambda a: (round(a["y"], 6), a["x"]))
                if direction == "right_to_left":
                    by_y = {}
                    for a in subzone_aps:
                        by_y.setdefault(round(a["y"], 6), []).append(a)
                    subzone_aps = []
                    for yk in sorted(by_y.keys()):
                        subzone_aps.extend(sorted(by_y[yk], key=lambda a: a["x"], reverse=True))
            else:
                subzone_aps.sort(key=lambda a: (round(a["x"], 6), a["y"]))
                if direction == "bottom_to_top":
                    by_x = {}
                    for a in subzone_aps:
                        by_x.setdefault(round(a["x"], 6), []).append(a)
                    subzone_aps = []
                    for xk in sorted(by_x.keys()):
                        subzone_aps.extend(sorted(by_x[xk], key=lambda a: a["y"], reverse=True))
        locations = []
        for idx, ap in enumerate(subzone_aps, 1):
            if entry_exit_type == "LM":
                ep = find_closest_point_by_path_distance(
                    ap["name"], lms, graph, node_positions, "LM")
            else:
                ep = find_closest_point_by_path_distance(
                    ap["name"], aps, graph, node_positions, "AP")
            locations.append({
                "customerLocation": f"{subzone_name}-{idx:03d}",
                "fleetManagerLocation": ap['name'],
                "locationType": "BOTH",
                "entryPoint": ep,
                "exitPoint": ep
            })
        subzones.append({
            "name": subzone_name,
            "entryPoint": sub_entry,
            "exitPoint": sub_exit,
            "locationSelectionStrategy": "SEQUENTIAL",
            "locations": locations
        })
    if current_zone:
        current_zone["subZones"] = subzones
        zones.append(current_zone)
    data["LocationDataMapping"]["zones"] = zones
    return data


# ─────────────────────────────────────────────
# N-Deep zone configuration helpers
# ─────────────────────────────────────────────

def _spatial_group(aps, primary='y', secondary='x', tolerance=0.3):
    sorted_aps = sorted(aps, key=lambda a: (a[primary], a[secondary]))
    groups, current = [], [sorted_aps[0]]
    ref = sorted_aps[0][primary]
    for ap in sorted_aps[1:]:
        if abs(ap[primary] - ref) <= tolerance:
            current.append(ap)
        else:
            groups.append(sorted(current, key=lambda a: a[secondary]))
            current = [ap]
            ref = ap[primary]
    groups.append(sorted(current, key=lambda a: a[secondary]))
    return groups


def _score_grouping(groups, graph):
    score = 0
    for group in groups:
        for i in range(len(group) - 1):
            a, b = group[i]["name"], group[i + 1]["name"]
            if b in graph.get(a, {}):
                score += 1
    return score


def _split_consecutive(sorted_aps, graph, zone_ap_names, factor=2.5):
    if len(sorted_aps) <= 1:
        return [sorted_aps]
    eucl_gaps = [
        math.hypot(sorted_aps[i + 1]["x"] - sorted_aps[i]["x"],
                   sorted_aps[i + 1]["y"] - sorted_aps[i]["y"])
        for i in range(len(sorted_aps) - 1)
    ]
    min_gap = min(eucl_gaps)
    gap_threshold = max(min_gap * 2.5, 0.5)
    subgroups, current = [], [sorted_aps[0]]
    for i, eucl in enumerate(eucl_gaps):
        prev_ap = sorted_aps[i]
        curr_ap = sorted_aps[i + 1]
        prev_name = prev_ap["name"]
        curr_name = curr_ap["name"]
        if curr_name in graph.get(prev_name, {}):
            current.append(curr_ap)
            continue
        if eucl > gap_threshold:
            subgroups.append(current)
            current = [curr_ap]
            continue
        excluded = zone_ap_names - {prev_name, curr_name}
        restricted = {
            n: {v: w for v, w in nb.items() if v not in excluded}
            for n, nb in graph.items()
            if n not in excluded
        }
        connected = False
        if prev_name in restricted:
            dists, _ = dijkstra(prev_name, restricted)
            path_d = dists.get(curr_name, float('inf'))
            if not math.isinf(path_d) and path_d <= factor * max(eucl, 0.01):
                connected = True
        if connected:
            current.append(curr_ap)
        else:
            subgroups.append(current)
            current = [curr_ap]
    subgroups.append(current)
    return subgroups


def _group_aps_into_inline_subzones(zone_aps, graph, tolerance=0.3):
    if not zone_aps:
        return []
    if len(zone_aps) == 1:
        return [zone_aps]
    y_groups = _spatial_group(zone_aps, primary='y', secondary='x', tolerance=tolerance)
    x_groups = _spatial_group(zone_aps, primary='x', secondary='y', tolerance=tolerance)
    y_score = _score_grouping(y_groups, graph)
    x_score = _score_grouping(x_groups, graph)
    raw_groups = y_groups if y_score >= x_score else x_groups
    zone_ap_names = {ap["name"] for ap in zone_aps}
    final_groups = []
    for group in raw_groups:
        subgroups = _split_consecutive(group, graph, zone_ap_names)
        final_groups.extend(subgroups)
    return [g for g in final_groups if g]


def _find_nearest_lm_to_ap(ap_name, lms_raw, graph, node_positions):
    lm_points = [r["_raw"] for r in lms_raw]
    return find_closest_point_by_path_distance(ap_name, lm_points, graph, node_positions, "LM")


def _order_subzone_aps(subzone_aps, zone_entry_name, graph, node_positions, drop_seq="f2l"):
    dists, _ = dijkstra(zone_entry_name, graph)

    def ap_key(ap):
        d = dists.get(ap["name"], float('inf'))
        if math.isinf(d):
            ref = node_positions.get(zone_entry_name)
            if ref:
                d = math.hypot(ap["x"] - ref[0], ap["y"] - ref[1])
        return d

    ordered = sorted(subzone_aps, key=ap_key)
    if drop_seq == "l2f":
        ordered = list(reversed(ordered))
    return ordered


def generate_ndeep_zone_configuration(mapping_data, df, drop_seq,
                                       zone_scope, sub_scope, loc_scope,
                                       entry_exit_type, graph, node_positions,
                                       log_fn=None):
    aps_raw, lms_raw = [], []
    for p in mapping_data.get("advancedPointList", []):
        cls = p.get("className")
        pos = p.get("pos", {})
        if "x" not in pos or "y" not in pos:
            continue
        rec = {
            "name": p.get("instanceName", ""),
            "x": float(pos["x"]),
            "y": float(pos["y"]),
            "_raw": p,
        }
        if cls == "ActionPoint":
            aps_raw.append(rec)
        elif cls == "LocationMark":
            lms_raw.append(rec)

    zones_out = []
    seen_zones = []
    for zone_name in df['Zone name']:
        if zone_name not in seen_zones:
            seen_zones.append(zone_name)

    for zone_name in seen_zones:
        zone_rows = df[df['Zone name'] == zone_name]
        row0 = zone_rows.iloc[0]
        zone_entry_name = str(row0['zone entry'])
        zone_exit_name = str(row0['zone exit'])
        min_x = float(zone_rows['min x'].min())
        max_x = float(zone_rows['max x'].max())
        min_y = float(zone_rows['min y'].min())
        max_y = float(zone_rows['max y'].max())

        zone_aps = [a for a in aps_raw
                    if min_x <= a["x"] <= max_x and min_y <= a["y"] <= max_y]
        if not zone_aps:
            if log_fn:
                log_fn(f"No APs found in zone '{zone_name}', skipping")
            continue

        if log_fn:
            log_fn(f"Zone '{zone_name}': {len(zone_aps)} APs found")

        subzone_groups = _group_aps_into_inline_subzones(zone_aps, graph)
        if log_fn:
            log_fn(f"  → {len(subzone_groups)} subzone(s) detected")

        if zone_entry_name in graph:
            dists_entry, _ = dijkstra(zone_entry_name, graph)
            subzone_groups.sort(
                key=lambda grp: min(
                    dists_entry.get(ap["name"], float('inf')) for ap in grp))

        subzones_out = []
        for sz_idx, sz_aps in enumerate(subzone_groups, 1):
            sz_name = f"{zone_name}_{sz_idx}"
            ordered_aps = _order_subzone_aps(sz_aps, zone_entry_name, graph, node_positions, drop_seq)
            ap_entry_side = ordered_aps[0]
            ap_exit_side = ordered_aps[-1]
            sz_entry_lm = _find_nearest_lm_to_ap(ap_entry_side["name"], lms_raw, graph, node_positions)
            sz_exit_lm = _find_nearest_lm_to_ap(ap_exit_side["name"], lms_raw, graph, node_positions)

            locations = []
            for loc_idx, ap in enumerate(ordered_aps, 1):
                fleet_loc = ap["name"]
                cust_loc = f"{sz_name}-{loc_idx:03d}"
                if entry_exit_type == "LM":
                    ep = _find_nearest_lm_to_ap(fleet_loc, lms_raw, graph, node_positions)
                    xp = ep
                else:
                    if loc_idx == 1:
                        ep = _find_nearest_lm_to_ap(fleet_loc, lms_raw, graph, node_positions)
                    else:
                        ep = ordered_aps[loc_idx - 2]["name"]
                    xp = ep

                loc_dict = {
                    "customerLocation": cust_loc,
                    "fleetManagerLocation": fleet_loc,
                    "locationType": "BOTH",
                }
                if loc_scope in ("BOTH", "ENTRY"):
                    loc_dict["entryPoint"] = ep
                if loc_scope in ("BOTH", "EXIT"):
                    loc_dict["exitPoint"] = xp
                locations.append(loc_dict)

            sz_dict = {
                "name": sz_name,
                "locationSelectionStrategy": "SUBZONELIFO",
                "locations": locations,
            }
            if sub_scope in ("BOTH", "ENTRY"):
                sz_dict["entryPoint"] = sz_entry_lm
            if sub_scope in ("BOTH", "EXIT"):
                sz_dict["exitPoint"] = sz_exit_lm
            subzones_out.append(sz_dict)

        zone_dict = {
            "name": zone_name,
            "locationSelectionStrategy": "SUBZONELIFO",
            "subZones": subzones_out,
        }
        if zone_scope in ("BOTH", "ENTRY"):
            zone_dict["entryPoint"] = zone_entry_name
        if zone_scope in ("BOTH", "EXIT"):
            zone_dict["exitPoint"] = zone_exit_name
        zones_out.append(zone_dict)

    return {"LocationDataMapping": {"zones": zones_out}}


# ─────────────────────────────────────────────
# Tab 4 — Map Validator
# ─────────────────────────────────────────────

def check_connectivity(all_points, outgoing):
    """
    Tarjan's iterative SCC + condensation BFS reachability.
    Returns (issue_rows, n_sccs).
    issue_rows: list of (src_label, src_type, issue_text, src_pts, dst_pts)
    """
    names = set(all_points.keys())
    flt_out = {v: outgoing[v] & names for v in names if v in outgoing}

    idx_map, lowlink, on_stk, stk, ctr, sccs = {}, {}, set(), [], [0], []

    def _visit(root):
        work = [(root, iter(flt_out.get(root, set())))]
        idx_map[root] = lowlink[root] = ctr[0]
        ctr[0] += 1
        stk.append(root)
        on_stk.add(root)
        while work:
            v, nbrs = work[-1]
            try:
                w = next(nbrs)
                if w not in idx_map:
                    idx_map[w] = lowlink[w] = ctr[0]
                    ctr[0] += 1
                    stk.append(w)
                    on_stk.add(w)
                    work.append((w, iter(flt_out.get(w, set()))))
                elif w in on_stk:
                    lowlink[v] = min(lowlink[v], idx_map[w])
            except StopIteration:
                work.pop()
                if work:
                    lowlink[work[-1][0]] = min(lowlink[work[-1][0]], lowlink[v])
                if lowlink[v] == idx_map[v]:
                    scc = []
                    while True:
                        w = stk.pop()
                        on_stk.discard(w)
                        scc.append(w)
                        if w == v:
                            break
                    sccs.append(frozenset(scc))

    for n in sorted(names):
        if n not in idx_map:
            _visit(n)

    if len(sccs) <= 1:
        return [], len(sccs)

    node_scc = {}
    for i, scc in enumerate(sccs):
        for n in scc:
            node_scc[n] = i

    n_s = len(sccs)
    cond = [set() for _ in range(n_s)]
    for v in names:
        sv = node_scc[v]
        for w in flt_out.get(v, set()):
            sw = node_scc.get(w)
            if sw is not None and sv != sw:
                cond[sv].add(sw)

    reachable = []
    for i in range(n_s):
        vis = set()
        q = deque([i])
        while q:
            cur = q.popleft()
            for nxt in cond[cur]:
                if nxt not in vis:
                    vis.add(nxt)
                    q.append(nxt)
        reachable.append(vis)

    rows = []
    for i in range(n_s):
        for j in range(n_s):
            if i == j or j in reachable[i]:
                continue
            src_pts = sorted(sccs[i])
            dst_pts = sorted(sccs[j])
            src_rep = src_pts[0]
            src_type = all_points.get(src_rep, "?")
            src_label = (src_rep if len(src_pts) == 1
                         else f"{src_rep} (grp {len(src_pts)})")
            preview = ", ".join(dst_pts[:4])
            if len(dst_pts) > 4:
                preview += f"  +{len(dst_pts) - 4} more"
            issue = f"Cannot reach {len(dst_pts)} pt(s): {preview}"
            rows.append((src_label, src_type, issue, src_pts, dst_pts))

    return rows, n_s


def prune_bidirectional_points(all_points, outgoing, incoming):
    """
    Removes points whose every edge is bidirectional (each neighbor has both
    an outgoing and incoming edge to/from the point — i.e. no one-way edge
    touches it at all). Points with no edges are left alone (they're already
    caught by the entry/exit check).

    Removing such a point can turn a former neighbor into an all-bidirectional
    point too (a chain of reversible-only points), so this repeats in rounds
    until nothing more qualifies.

    Returns (pruned_points, pruned_outgoing, pruned_incoming, removed) where
    removed is a sorted list of (name, type) tuples for every point taken out.
    """
    pts = dict(all_points)
    out = {k: set(v) for k, v in outgoing.items()}
    inc = {k: set(v) for k, v in incoming.items()}
    removed = []

    while True:
        candidates = [name for name in pts if out.get(name) and out.get(name) == inc.get(name)]
        if not candidates:
            break
        for name in candidates:
            removed.append((name, pts[name]))
            for n in out.get(name, set()) | inc.get(name, set()):
                out.get(n, set()).discard(name)
                inc.get(n, set()).discard(name)
            out.pop(name, None)
            inc.pop(name, None)
            pts.pop(name, None)

    return pts, out, inc, sorted(removed)


def run_map_validation(json_data, detailed_review=False):
    """
    Full map validation. Returns dict with issues, connectivity rows, counts.
    This is CPU-bound — the caller should run it in a thread.

    When detailed_review is True, points that are only ever connected by
    bidirectional edges are pruned out first (see prune_bidirectional_points)
    and the checks below run on that reduced graph. When False (default),
    the checks run on the map exactly as uploaded.
    """
    points = json_data.get("advancedPointList", [])
    curves = json_data.get("advancedCurveList", [])

    outgoing = defaultdict(set)
    incoming = defaultdict(set)
    for curve in curves:
        src = (curve.get("startPos") or {}).get("instanceName")
        dst = (curve.get("endPos") or {}).get("instanceName")
        if src and dst and src != dst:
            outgoing[src].add(dst)
            incoming[dst].add(src)

    all_points = {}
    for pt in points:
        name = pt.get("instanceName")
        cls = pt.get("className", "Unknown")
        if name:
            all_points[name] = "AP" if cls == "ActionPoint" else "LM"

    if detailed_review:
        pruned_points, pruned_out, pruned_in, removed = prune_bidirectional_points(
            all_points, outgoing, incoming
        )
    else:
        pruned_points, pruned_out, pruned_in, removed = all_points, outgoing, incoming, []

    issues = []
    for name, pt_type in sorted(pruned_points.items()):
        has_in = bool(pruned_in.get(name))
        has_out = bool(pruned_out.get(name))
        if not has_in and not has_out:
            issues.append((name, pt_type, "No entry path  &  No exit path"))
        elif not has_in:
            issues.append((name, pt_type, "No entry path"))
        elif not has_out:
            issues.append((name, pt_type, "No exit path"))

    conn_rows, n_sccs = check_connectivity(pruned_points, pruned_out)

    return {
        "all_points": pruned_points,
        "issues": issues,
        "conn_rows": [
            {
                "src_label": r[0], "src_type": r[1], "issue": r[2],
                "src_pts": r[3], "dst_pts": r[4]
            }
            for r in conn_rows
        ],
        "n_sccs": n_sccs,
        "total": len(pruned_points),
        "n_issues": len(issues),
        "n_conn": len(conn_rows),
        "removed_points": [{"name": n, "type": t} for n, t in removed],
        "n_removed": len(removed),
        "total_before": len(all_points),
        "detailed_review": detailed_review,
    }


# ─────────────────────────────────────────────
# Sample Excel builder (Point to Point tab)
# ─────────────────────────────────────────────

def build_sample_excel(output_path):
    wb = Workbook()

    # ── Sheet 1: Standard ──
    ws1 = wb.active
    ws1.title = "Standard"
    std_headers = ["Zone name", "zone entry", "zone exit",
                   "subzone name", "subzone entery", "subzone exit",
                   "min x", "max x", "min y", "max y"]
    std_rows = [
        ["Zone_A", "LM_Entry_A", "LM_Exit_A",
         "SubZone_A1", "LM_SubEntry_A1", "LM_SubExit_A1", 0.0, 10.0, 0.0, 5.0],
        ["Zone_A", "LM_Entry_A", "LM_Exit_A",
         "SubZone_A2", "LM_SubEntry_A2", "LM_SubExit_A2", 0.0, 10.0, 5.1, 10.0],
        ["Zone_B", "LM_Entry_B", "LM_Exit_B",
         "SubZone_B1", "LM_SubEntry_B1", "LM_SubExit_B1", 11.0, 20.0, 0.0, 10.0],
    ]

    hdr_fill = PatternFill("solid", fgColor="494BD6")
    alt_fill = PatternFill("solid", fgColor="1E2435")
    ws1.append(std_headers)
    for cell in ws1[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(std_rows):
        ws1.append(row)
        if i % 2 == 1:
            for cell in ws1[i + 2]:
                cell.fill = alt_fill
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = max(
            len(str(col[0].value or "")), 14) + 2

    # ── Sheet 2: N-Deep ──
    ws2 = wb.create_sheet("N-Deep")
    nd_headers = ["Zone name", "zone entry", "zone exit",
                  "min x", "max x", "min y", "max y"]
    nd_rows = [
        ["Zone_A", "LM_Entry_A", "LM_Exit_A", 0.0, 10.0, 0.0, 10.0],
        ["Zone_B", "LM_Entry_B", "LM_Exit_B", 11.0, 20.0, 0.0, 10.0],
    ]
    ws2.append(nd_headers)
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(nd_rows):
        ws2.append(row)
        if i % 2 == 1:
            for cell in ws2[i + 2]:
                cell.fill = alt_fill
    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = max(
            len(str(col[0].value or "")), 14) + 2

    wb.save(output_path)


# ─────────────────────────────────────────────
# Tab 5 — Orderfile Generator
# ─────────────────────────────────────────────

class OrderfileOptimizer:
    def __init__(self, df, target_orders, target_lines, target_quantity,
                 max_qty_per_order=None, min_qty_per_order=None, max_qty_per_line=None):
        self.original_df = df.copy()
        self.target_orders = target_orders
        self.target_lines = target_lines
        self.target_quantity = target_quantity
        self.max_qty_per_order = max_qty_per_order
        self.min_qty_per_order = min_qty_per_order
        self.max_qty_per_line = max_qty_per_line

        if max_qty_per_line is not None:
            df = df.copy()
            df['unit'] = df['unit'].clip(upper=max_qty_per_line)

        self.df = df

        self.order_stats = df.groupby('order').agg(
            lines=('line', 'sum'),
            total_qty=('unit', 'sum')
        ).reset_index()

        if min_qty_per_order is not None:
            self.order_stats = self.order_stats[
                self.order_stats['total_qty'] >= min_qty_per_order
            ].reset_index(drop=True)
            valid_orders = self.order_stats['order'].values
            self.df = self.df[self.df['order'].isin(valid_orders)].copy()

        if max_qty_per_order is not None:
            self.order_stats['total_qty'] = self.order_stats['total_qty'].clip(
                upper=max_qty_per_order)

        self.orders_list = self.order_stats['order'].values
        self.lines_list = self.order_stats['lines'].values.astype(float)
        self.qty_list = self.order_stats['total_qty'].values.astype(float)
        self.total_orders = len(self.orders_list)

        self.avg_lines = self.lines_list.mean() if self.total_orders > 0 else 1
        self.avg_qty = self.qty_list.mean() if self.total_orders > 0 else 1

        self.sorted_by_lines = np.argsort(self.lines_list)
        self.sorted_by_qty = np.argsort(self.qty_list)

    def calculate_score(self, solution):
        selected = solution.astype(bool)
        orders = float(np.sum(selected))
        lines = float(np.sum(self.lines_list[selected]))
        qty = float(np.sum(self.qty_list[selected]))

        order_pct_err = abs(orders - self.target_orders) / max(self.target_orders, 1)
        line_pct_err  = abs(lines  - self.target_lines)  / max(self.target_lines, 1)
        qty_pct_err   = abs(qty    - self.target_quantity)/ max(self.target_quantity, 1)

        order_over = max(0, orders - self.target_orders) / max(self.target_orders, 1)
        line_over  = max(0, lines  - self.target_lines)  / max(self.target_lines, 1)
        qty_over   = max(0, qty    - self.target_quantity)/ max(self.target_quantity, 1)

        return (order_pct_err * 100 + line_pct_err * 100 + qty_pct_err * 100
                + (order_over + line_over + qty_over) * 200)

    def _get_totals(self, solution):
        sel = solution.astype(bool)
        return (float(np.sum(sel)),
                float(np.sum(self.lines_list[sel])),
                float(np.sum(self.qty_list[sel])))

    def generate_initial_solution(self, strategy='balanced'):
        solution = np.zeros(self.total_orders, dtype=bool)
        if strategy == 'balanced':
            target_avg_lines = self.target_lines / max(self.target_orders, 1)
            target_avg_qty   = self.target_quantity / max(self.target_orders, 1)
            line_fit = np.abs(self.lines_list - target_avg_lines) / max(target_avg_lines, 1)
            qty_fit  = np.abs(self.qty_list   - target_avg_qty)   / max(target_avg_qty, 1)
            sorted_idx = np.argsort(line_fit + qty_fit)
        elif strategy == 'lines_first':
            target_avg_lines = self.target_lines / max(self.target_orders, 1)
            sorted_idx = np.argsort(np.abs(self.lines_list - target_avg_lines))
        elif strategy == 'qty_first':
            target_avg_qty = self.target_quantity / max(self.target_orders, 1)
            sorted_idx = np.argsort(np.abs(self.qty_list - target_avg_qty))
        elif strategy == 'random':
            sorted_idx = np.random.permutation(self.total_orders)
        else:
            sorted_idx = np.arange(self.total_orders)

        count = 0
        for i in sorted_idx:
            if count >= self.target_orders:
                break
            solution[i] = True
            count += 1
        return solution

    def generate_neighbor(self, sol):
        neighbor = sol.copy()
        selected   = np.where(neighbor)[0]
        unselected = np.where(~neighbor)[0]

        if len(selected) == 0 or len(unselected) == 0:
            idx = random.randint(0, self.total_orders - 1)
            neighbor[idx] = not neighbor[idx]
            return neighbor

        orders, lines, qty = self._get_totals(neighbor)
        order_err = orders - self.target_orders
        line_err  = lines  - self.target_lines
        qty_err   = qty    - self.target_quantity

        r = random.random()

        if r < 0.40:
            worst = max(
                ('order', abs(order_err) / max(self.target_orders, 1)),
                ('line',  abs(line_err)  / max(self.target_lines, 1)),
                ('qty',   abs(qty_err)   / max(self.target_quantity, 1)),
                key=lambda x: x[1]
            )[0]
            remove_idx = random.choice(selected)
            if worst == 'line':
                needed = self.lines_list[remove_idx] - line_err
                diffs = np.abs(self.lines_list[unselected] - needed)
                add_idx = random.choice(unselected[np.argsort(diffs)[:5]])
            elif worst == 'qty':
                needed = self.qty_list[remove_idx] - qty_err
                diffs = np.abs(self.qty_list[unselected] - needed)
                add_idx = random.choice(unselected[np.argsort(diffs)[:5]])
            else:
                add_idx = random.choice(unselected)
            neighbor[remove_idx] = False
            neighbor[add_idx] = True

        elif r < 0.60:
            if order_err > 0 and len(selected) > 0:
                excess_lines = max(0, line_err)
                excess_qty   = max(0, qty_err)
                if excess_lines > 0 or excess_qty > 0:
                    scores = (self.lines_list[selected] * (1 if line_err > 0 else -1) +
                              self.qty_list[selected]   * (1 if qty_err  > 0 else -1))
                    neighbor[random.choice(selected[np.argsort(-scores)[:3]])] = False
                else:
                    neighbor[random.choice(selected)] = False
            elif order_err < 0 and len(unselected) > 0:
                deficit_lines = max(0, -line_err)
                deficit_qty   = max(0, -qty_err)
                if deficit_lines > 0 or deficit_qty > 0:
                    scores = (np.minimum(self.lines_list[unselected], deficit_lines) +
                              np.minimum(self.qty_list[unselected],   deficit_qty))
                    neighbor[random.choice(unselected[np.argsort(-scores)[:3]])] = True
                else:
                    small = unselected[np.argsort(
                        self.lines_list[unselected] + self.qty_list[unselected])[:3]]
                    neighbor[random.choice(small)] = True
            else:
                neighbor[random.choice(selected)] = False
                neighbor[random.choice(unselected)] = True

        elif r < 0.80:
            n_swap = min(2, len(selected), len(unselected))
            for idx in np.random.choice(selected,   n_swap, replace=False): neighbor[idx] = False
            for idx in np.random.choice(unselected, n_swap, replace=False): neighbor[idx] = True

        else:
            n_flip = random.randint(1, min(3, self.total_orders))
            for idx in np.random.choice(self.total_orders, n_flip, replace=False):
                neighbor[idx] = not neighbor[idx]

        return neighbor

    def simulated_annealing(self, sol, time_limit=60):
        current = sol.copy()
        current_score = self.calculate_score(current)
        best = current.copy()
        best_score = current_score

        T, T_min, cooling = 50.0, 0.01, 0.995
        stagnant = 0
        start = time.time()

        while T > T_min and time.time() - start < time_limit:
            improved_this_round = False
            for _ in range(100):
                if time.time() - start >= time_limit:
                    break
                neighbor = self.generate_neighbor(current)
                score = self.calculate_score(neighbor)
                delta = score - current_score
                if delta < 0 or random.random() < math.exp(-delta / max(T, 0.001)):
                    current = neighbor
                    current_score = score
                    if score < best_score:
                        best = neighbor.copy()
                        best_score = score
                        improved_this_round = True
            stagnant = 0 if improved_this_round else stagnant + 1
            if stagnant > 20:
                T = min(T * 5, 50.0)
                stagnant = 0
            T *= cooling

        return best

    def local_search(self, sol, time_limit=30):
        best = sol.copy()
        best_score = self.calculate_score(best)
        start = time.time()

        improved = True
        while improved and time.time() - start < time_limit * 0.5:
            improved = False
            for i in np.random.permutation(self.total_orders):
                if time.time() - start >= time_limit * 0.5:
                    break
                test = best.copy()
                test[i] = not test[i]
                score = self.calculate_score(test)
                if score < best_score:
                    best = test
                    best_score = score
                    improved = True

        selected   = np.where(best)[0]
        unselected = np.where(~best)[0]
        if len(selected) > 0 and len(unselected) > 0:
            improved = True
            while improved and time.time() - start < time_limit:
                improved = False
                np.random.shuffle(selected)
                for rem in selected[:min(len(selected), 50)]:
                    if time.time() - start >= time_limit:
                        break
                    for add in np.random.choice(unselected, min(len(unselected), 30), replace=False):
                        test = best.copy()
                        test[rem] = False
                        test[add] = True
                        score = self.calculate_score(test)
                        if score < best_score:
                            best = test
                            best_score = score
                            selected   = np.where(best)[0]
                            unselected = np.where(~best)[0]
                            improved = True
                            break

        return best

    def scale_quantities(self, result_df, target_qty):
        if result_df.empty or target_qty <= 0:
            return result_df
        current_total = result_df['unit'].sum()
        if current_total == 0:
            return result_df

        result_df = result_df.copy()
        scale_factor = target_qty / current_total
        ideal = result_df['unit'].values * scale_factor
        floored = np.maximum(np.floor(ideal).astype(int), 1)

        if self.max_qty_per_line is not None:
            floored = np.minimum(floored, int(self.max_qty_per_line))

        remainder = int(target_qty) - int(np.sum(floored))
        if remainder > 0:
            candidates = np.argsort(-(ideal - floored))
            for idx in candidates:
                if remainder <= 0:
                    break
                cap = int(self.max_qty_per_line) if self.max_qty_per_line else float('inf')
                if floored[idx] < cap:
                    floored[idx] += 1
                    remainder -= 1
        elif remainder < 0:
            for idx in np.argsort(-floored):
                if remainder >= 0:
                    break
                if floored[idx] > 1:
                    floored[idx] -= 1
                    remainder += 1

        result_df['unit'] = floored

        if self.max_qty_per_order is not None:
            max_oq = int(self.max_qty_per_order)
            for order_id, grp in result_df.groupby('order'):
                order_total = grp['unit'].sum()
                if order_total > max_oq:
                    excess = order_total - max_oq
                    for idx in grp.sort_values('unit', ascending=False).index:
                        if excess <= 0:
                            break
                        reduce = min(excess, result_df.loc[idx, 'unit'] - 1)
                        if reduce > 0:
                            result_df.loc[idx, 'unit'] -= reduce
                            excess -= reduce

        if self.min_qty_per_order is not None:
            min_oq = int(self.min_qty_per_order)
            for order_id, grp in result_df.groupby('order'):
                order_total = grp['unit'].sum()
                if order_total < min_oq:
                    deficit = min_oq - order_total
                    idxs = grp.index
                    per_line = max(1, deficit // len(idxs))
                    for idx in idxs:
                        if deficit <= 0:
                            break
                        add = min(per_line, deficit)
                        cap = int(self.max_qty_per_line) if self.max_qty_per_line else float('inf')
                        add = min(add, cap - result_df.loc[idx, 'unit'])
                        if add > 0:
                            result_df.loc[idx, 'unit'] += add
                            deficit -= add

        return result_df

    def optimize(self, progress_cb=None, total_time=120):
        strategies = ['balanced', 'lines_first', 'qty_first', 'random']
        n_restarts = len(strategies)
        time_per_restart = total_time / n_restarts

        global_best_score = float('inf')
        global_best_sol = None

        for i, strategy in enumerate(strategies):
            pct_base = int(i / n_restarts * 100)
            if progress_cb:
                progress_cb(min(pct_base + 2, 99),
                            f"Restart {i+1}/{n_restarts} ({strategy})…")

            sol = self.generate_initial_solution(strategy=strategy)
            sol = self.simulated_annealing(sol, time_per_restart * 0.55)

            if progress_cb:
                progress_cb(min(pct_base + int(60 / n_restarts), 99),
                            f"Restart {i+1}/{n_restarts} — local search…")

            sol = self.local_search(sol, time_per_restart * 0.45)

            score = self.calculate_score(sol)
            if score < global_best_score:
                global_best_score = score
                global_best_sol = sol.copy()

        if progress_cb:
            progress_cb(90, "Final polish…")

        if global_best_sol is not None:
            global_best_sol = self.local_search(global_best_sol, total_time * 0.05)

        if progress_cb:
            progress_cb(95, "Scaling quantities to hit target…")

        sol = global_best_sol
        selected_orders = self.orders_list[sol]
        result = self.df[self.df["order"].isin(selected_orders)].copy()
        result = self.scale_quantities(result, self.target_quantity)

        if progress_cb:
            progress_cb(100, "Done!")

        actual_orders = len(selected_orders)
        actual_lines  = int(np.sum(self.lines_list[sol]))
        actual_qty    = float(result['unit'].sum())

        return result, actual_orders, actual_lines, actual_qty


def orderfile_to_excel(df, output_path):
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Orders')
